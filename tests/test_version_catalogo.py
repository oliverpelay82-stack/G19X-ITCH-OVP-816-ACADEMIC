"""
M6 — Versión del catálogo de reglas en el core (PRD §21.6; protocolo de
auditoría del DOC-05; HU-54).

  1. El cliente del rules-service toma la versión de la cabecera
     X-Version-Catalogo del inventario y la evaluación del expediente la informa.
  2. El dictamen guarda version_catalogo, la sella en su evento y el trigger de
     RN-05 impide cambiarla; un dictamen sin ella la deja vacía.
  3. El Excel del expediente la muestra en el resumen y en la hoja «Evidencia».

Aislamiento: base de datos de pruebas plurione_compliance_test (tests/bd_pruebas.py).
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bd_pruebas  # noqa: E402

bd_pruebas.preparar_entorno()
os.environ["REEVALUACION_INTERVALO_HORAS"] = "0"

from io import BytesIO  # noqa: E402
from types import SimpleNamespace  # noqa: E402

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import text  # noqa: E402
from sqlalchemy.exc import DBAPIError  # noqa: E402
from main import app  # noqa: E402
from clients import rules_client  # noqa: E402
from config.database import SessionLocal, engine, init_db  # noqa: E402
from models.titular import TitularExpediente  # noqa: E402

bd_pruebas.verificar_aislamiento(engine)
init_db()
VERSION = "cat-0123456789ab"
_obtener_real = rules_client.obtener_documentos_obligatorios
VERSIONES_TIPO = {"*": "tip-general00000", "CSF_SAT": "tip-csf000000001"}
rules_client.obtener_documentos_obligatorios = lambda *_: rules_client.InventarioObligatorio(
    documentos={"CSF_SAT": {"codigo_regla": "RIE-01", "severidad": "CRITICA"}},
    fuente=rules_client.FUENTE_CATALOGO, version_catalogo=VERSION, versiones_por_tipo=dict(VERSIONES_TIPO),
)
auth = bd_pruebas.Autenticacion()
client = auth.cliente(app)


def _documento(rfc: str) -> dict:
    db = SessionLocal()
    try:
        titular = TitularExpediente(rfc=rfc, razon_social=f"Catálogo {rfc} S.A. de C.V.", tipo_persona="MORAL")
        db.add(titular)
        db.commit()
        titular_id = titular.id
    finally:
        db.close()
    exp_id = client.post("/api/v1/expedientes", json={"titular_id": titular_id}).json()["id"]
    r = client.post("/api/v1/documentos/upload", data={"expediente_id": exp_id, "tipo_documento": "CSF_SAT"},
                    files={"file": ("csf.pdf", b"%PDF-1.4 catalogo " + os.urandom(8), "application/pdf")})
    assert r.status_code == 201, r.text
    return r.json()


def _dictaminar(doc: dict, **extra) -> dict:
    r = client.post("/api/v1/dictamenes", headers=auth.sistema(), json={"documento_id": doc["id"], "estatus": "EN_REVISION",
                                                 "nivel_confianza": 0.8, "observaciones": ["M6."], **extra})
    assert r.status_code == 201, r.text
    return r.json()


def test_el_cliente_toma_la_version_de_la_cabecera(monkeypatch):
    import requests

    respuesta = SimpleNamespace(
        json=lambda: [{"tipo_documento": "CSF_SAT", "codigo_regla": "RIE-01-CSF", "severidad": "CRITICA"}],
        headers={"X-Version-Catalogo": VERSION}, raise_for_status=lambda: None)
    monkeypatch.setattr(requests, "get", lambda *a, **k: respuesta)
    assert _obtener_real("CORPORATIVO").version_catalogo == VERSION


def test_la_evaluacion_informa_la_version():
    doc = _documento("VCT010101AAA")
    evaluacion = client.get(f"/api/v1/expedientes/{doc['expediente_id']}").json()["evaluacion"]
    assert evaluacion["version_catalogo"] == VERSION


def test_el_dictamen_guarda_y_protege_la_version():
    doc = _documento("VCT020202AAA")
    dictamen = _dictaminar(doc, version_catalogo=VERSION)

    guardado = client.get("/api/v1/dictamenes", params={"documento_id": doc["id"]}).json()[0]
    assert guardado["version_catalogo"] == VERSION
    [evento] = client.get("/api/v1/bitacora", params={"documento_id": doc["id"], "evento": "DICTAMEN_IA"}).json()
    assert evento["detalle"]["version_catalogo"] == VERSION
    with pytest.raises(DBAPIError), engine.begin() as conexion:
        conexion.execute(text("UPDATE dictamenes_auditoria SET version_catalogo = 'cat-otro' WHERE id = :id"),
                         {"id": dictamen["id"]})

    sin_version = _documento("VCT030303AAA")
    _dictaminar(sin_version)
    assert client.get("/api/v1/dictamenes", params={"documento_id": sin_version["id"]}).json()[0]["version_catalogo"] is None


def test_el_dictamen_guarda_y_protege_el_nombre_del_documento_real():
    """PRD §21.9: el nombre que dio la IA a un archivo de otro tipo se guarda, se sella y no se modifica."""
    doc = _documento("VCT080808AAA")
    dictamen = _dictaminar(doc, tipo_documento_detectado="IDENTIFICACION_OFICIAL",
                           nombre_documento_real="Credencial para Votar INE")
    guardado = client.get("/api/v1/dictamenes", params={"documento_id": doc["id"]}).json()[0]
    assert guardado["nombre_documento_real"] == "Credencial para Votar INE"
    [evento] = client.get("/api/v1/bitacora", params={"documento_id": doc["id"], "evento": "DICTAMEN_IA"}).json()
    assert evento["detalle"]["nombre_documento_real"] == "Credencial para Votar INE"
    with pytest.raises(DBAPIError), engine.begin() as conexion:
        conexion.execute(text("UPDATE dictamenes_auditoria SET nombre_documento_real = 'Otro' WHERE id = :id"),
                         {"id": dictamen["id"]})
    # Sin discrepancia el campo queda vacío y el evento no lo lleva.
    otro = _documento("VCT090909AAA")
    _dictaminar(otro)
    assert client.get("/api/v1/dictamenes", params={"documento_id": otro["id"]}).json()[0]["nombre_documento_real"] is None
    [evento] = client.get("/api/v1/bitacora", params={"documento_id": otro["id"], "evento": "DICTAMEN_IA"}).json()
    assert "nombre_documento_real" not in evento["detalle"]


def test_el_excel_muestra_la_version():
    doc = _documento("VCT040404AAA")
    _dictaminar(doc, version_catalogo=VERSION)

    r = client.get(f"/api/v1/reportes/expedientes/{doc['expediente_id']}", params={"formato": "xlsx"})

    assert r.status_code == 200
    libro = load_workbook(BytesIO(r.content))
    evidencia = libro["Evidencia"]
    fila = dict(zip([c.value for c in evidencia[1]], [c.value for c in evidencia[2]]))
    assert fila["Catálogo de reglas"] == VERSION
    valores = {c.value for hoja in libro.worksheets for fila in hoja.iter_rows() for c in fila}
    assert "Versión del catálogo de reglas" in valores


# ── M7: aviso de nueva auditoría ──────────────────────────────────────


def _alertas_nueva_auditoria(doc: dict) -> list[dict]:
    evaluacion = client.get(f"/api/v1/expedientes/{doc['expediente_id']}").json()["evaluacion"]
    return [a for a in evaluacion["alertas"] if a["codigo"] == "REQUIERE_NUEVA_AUDITORIA"]


def test_con_las_mismas_reglas_no_hay_aviso():
    doc = _documento("VCT050505AAA")
    _dictaminar(doc, version_reglas_tipo=VERSIONES_TIPO["CSF_SAT"])
    assert not _alertas_nueva_auditoria(doc)


def test_si_cambiaron_las_reglas_del_tipo_se_pide_nueva_auditoria():
    doc = _documento("VCT060606AAA")
    _dictaminar(doc, version_reglas_tipo="tip-anterior0000")
    [alerta] = _alertas_nueva_auditoria(doc)
    assert alerta["severidad"] == "MEDIA" and "cambiaron después de auditar" in alerta["detalle"]
    evaluacion = client.get(f"/api/v1/expedientes/{doc['expediente_id']}").json()["evaluacion"]
    assert any(d["requiere_nueva_auditoria"] for d in evaluacion["documentos"])

    _dictaminar(doc, version_reglas_tipo=VERSIONES_TIPO["CSF_SAT"])   # se audita de nuevo
    assert not _alertas_nueva_auditoria(doc)


def test_un_dictamen_previo_sin_version_no_avisa():
    doc = _documento("VCT070707AAA")
    _dictaminar(doc)
    assert not _alertas_nueva_auditoria(doc)


def test_la_cabecera_por_tipo_se_interpreta_y_tolera_basura():
    assert rules_client._versiones_por_tipo('{"*":"tip-a","csf_sat":"tip-b"}') == {"*": "tip-a", "CSF_SAT": "tip-b"}
    assert rules_client._versiones_por_tipo("no es json") == {} and rules_client._versiones_por_tipo(None) == {}
    inventario = rules_client.InventarioObligatorio(documentos={}, fuente="x", versiones_por_tipo={"*": "tip-a"})
    assert inventario.version_de_tipo("PODER_NOTARIAL") == "tip-a"

