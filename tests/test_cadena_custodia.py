"""
P3 — Cadena de custodia sellada de la bitácora (PRD §21.3: PLUS-04, PLUS-05,
HU-39, HU-40).

  1. Cada evento queda sellado y enlazado con el anterior; la API expone el
     sello y la verificación de una cadena sin tocar es íntegra.
  2. Una alteración con los triggers desactivados, un evento borrado y un
     evento insertado sin sello por fuera del motor se detectan con su
     secuencia. (Cada prueba restaura la bitácora al terminar.)
  3. El génesis resume el historial previo al sellado: alterar ese historial
     también se detecta.
  4. Cargas concurrentes nunca bifurcan la cadena.
  5. La ficha de custodia (PLUS-06, HU-41), en PDF y en Excel, reúne huella,
     versiones, dictámenes, ratificaciones, eventos sellados y la verificación;
     reporta un documento alterado y su exportación queda sellada.

Aislamiento: base de datos de pruebas plurione_compliance_test (tests/bd_pruebas.py).
"""
import json
import os
import sys
import threading
from contextlib import contextmanager
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bd_pruebas  # noqa: E402

BOVEDA = bd_pruebas.preparar_entorno()
os.environ["REEVALUACION_INTERVALO_HORAS"] = "0"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from sqlalchemy import text  # noqa: E402
from main import app  # noqa: E402
from clients import rules_client  # noqa: E402
from config.database import SessionLocal, engine, init_db  # noqa: E402
from models.bitacora import BitacoraAuditoria  # noqa: E402
from models.titular import TitularExpediente  # noqa: E402
from motor import bitacora  # noqa: E402

bd_pruebas.verificar_aislamiento(engine)
init_db()
rules_client.obtener_documentos_obligatorios = lambda *_: rules_client.InventarioObligatorio(
    documentos={"CSF_SAT": {"codigo_regla": "RIE-01", "severidad": "CRITICA"}},
    fuente=rules_client.FUENTE_CATALOGO,
)
client = TestClient(app)


@contextmanager
def sin_triggers():
    """Simula a quien administra la base: desactiva la inmutabilidad (y la restaura)."""
    with engine.begin() as c:
        c.execute(text("ALTER TABLE bitacora_auditoria DISABLE TRIGGER trg_bitacora_inmutable"))
    try:
        with engine.begin() as c:
            yield c
    finally:
        with engine.begin() as c:
            c.execute(text("ALTER TABLE bitacora_auditoria ENABLE TRIGGER trg_bitacora_inmutable"))


def _registrar(n: int = 3, evento: str = "PRUEBA_CADENA_P3") -> list[str]:
    """Eventos propios de esta suite: no alteran los conteos de las demás (base compartida)."""
    db = SessionLocal()
    try:
        ids = [bitacora.registrar(db, evento, usuario="pruebas.p3", detalle={"n": i, "monto": 0.1 * i, "txt": "ñandú"}).id
               for i in range(n)]
        db.commit()
        return ids
    finally:
        db.close()


def _verificar() -> dict:
    r = client.get("/api/v1/bitacora/verificar-cadena")
    assert r.status_code == 200, r.text
    return r.json()


def _evento(evento_id: str) -> BitacoraAuditoria:
    db = SessionLocal()
    try:
        return db.get(BitacoraAuditoria, evento_id)
    finally:
        db.close()


def test_cada_evento_queda_sellado_y_enlazado():
    ids = _registrar(3)

    eventos = [_evento(i) for i in ids]
    assert [e.secuencia for e in eventos] == [eventos[0].secuencia + k for k in range(3)]
    assert eventos[1].hash_anterior == eventos[0].hash_evento and eventos[2].hash_anterior == eventos[1].hash_evento
    assert all(bitacora.huella_evento(e) == e.hash_evento for e in eventos), "El valor guardado es el sellado"
    resultado = _verificar()
    assert resultado["integra"] and resultado["sellada"] and resultado["ultimo_hash"] == eventos[-1].hash_evento
    listado = client.get("/api/v1/bitacora", params={"limite": 5}).json()
    assert {"secuencia", "hash_anterior", "hash_evento"} <= set(listado[0])
    db = SessionLocal()
    try:
        assert db.query(BitacoraAuditoria).filter(BitacoraAuditoria.evento == "SELLO_GENESIS").count() == 1
    finally:
        db.close()


def test_una_alteracion_con_los_triggers_desactivados_se_detecta():
    [evento_id] = _registrar(1)
    original = _evento(evento_id)

    with sin_triggers() as c:
        c.execute(text("UPDATE bitacora_auditoria SET usuario = 'otra.persona' WHERE id = :id"), {"id": evento_id})
    try:
        resultado = _verificar()
        assert not resultado["integra"]
        assert any(h["secuencia"] == original.secuencia and "alterado" in h["problema"] for h in resultado["hallazgos"])
    finally:
        with sin_triggers() as c:
            c.execute(text("UPDATE bitacora_auditoria SET usuario = :u WHERE id = :id"),
                      {"u": original.usuario, "id": evento_id})
    assert _verificar()["integra"], "Restaurado el valor, la cadena vuelve a ser íntegra"


def test_un_evento_borrado_se_detecta():
    ids = _registrar(3)
    medio = _evento(ids[1])
    columnas = ("id, evento, expediente_id, documento_id, dictamen_id, usuario, ip_origen, detalle, creado_en, "
                "secuencia, hash_anterior, hash_evento")

    with sin_triggers() as c:
        fila = c.execute(text(f"SELECT {columnas} FROM bitacora_auditoria WHERE id = :id"), {"id": medio.id}).mappings().one()
        c.execute(text("DELETE FROM bitacora_auditoria WHERE id = :id"), {"id": medio.id})
    try:
        problemas = [h["problema"] for h in _verificar()["hallazgos"] if h["secuencia"] == medio.secuencia + 1]
        assert any("borrados" in p for p in problemas) and any("enlace" in p for p in problemas), problemas
    finally:
        with engine.begin() as c:   # insertar sí está permitido: se repone la fila tal cual
            c.execute(text(f"INSERT INTO bitacora_auditoria ({columnas}) VALUES (:id, :evento, :expediente_id, "
                           ":documento_id, :dictamen_id, :usuario, :ip_origen, CAST(:detalle AS JSON), :creado_en, "
                           ":secuencia, :hash_anterior, :hash_evento)"),
                      {**fila, "detalle": json.dumps(fila["detalle"])})
    assert _verificar()["integra"]


def test_un_evento_insertado_sin_sello_se_detecta():
    _registrar(1)
    with engine.begin() as c:
        c.execute(text("INSERT INTO bitacora_auditoria (id, evento, usuario, detalle, creado_en) "
                       "VALUES ('intruso-p3', 'RATIFICACION_HUMANA', 'nadie', '{}', clock_timestamp()::timestamp)"))
    try:
        resultado = _verificar()
        assert not resultado["integra"]
        assert any(h["evento_id"] == "intruso-p3" and "por fuera del motor" in h["problema"]
                   for h in resultado["hallazgos"])
    finally:
        with sin_triggers() as c:
            c.execute(text("DELETE FROM bitacora_auditoria WHERE id = 'intruso-p3'"))
    assert _verificar()["integra"]


def test_el_genesis_resume_y_protege_el_historial_previo():
    _registrar(2)
    # Simula una base anterior a P3: ningún evento sellado.
    with sin_triggers() as c:
        c.execute(text("UPDATE bitacora_auditoria SET secuencia = NULL, hash_anterior = NULL, hash_evento = NULL"))
        previos = c.execute(text("SELECT count(*) FROM bitacora_auditoria")).scalar()

    [primero] = _registrar(1)

    db = SessionLocal()
    try:
        genesis = db.query(BitacoraAuditoria).filter(BitacoraAuditoria.secuencia == 1).one()
    finally:
        db.close()
    assert genesis.evento == "SELLO_GENESIS" and genesis.detalle["eventos_previos"] == previos
    assert _evento(primero).secuencia == 2 and _evento(primero).hash_anterior == genesis.hash_evento
    resultado = _verificar()
    assert resultado["integra"] and resultado["eventos_previos"] == previos

    with sin_triggers() as c:
        viejo = c.execute(text("SELECT id, usuario FROM bitacora_auditoria WHERE secuencia IS NULL "
                               "ORDER BY creado_en LIMIT 1")).one()
        c.execute(text("UPDATE bitacora_auditoria SET usuario = 'otra.persona' WHERE id = :id"), {"id": viejo.id})
    try:
        assert any("historial anterior" in h["problema"] for h in _verificar()["hallazgos"])
    finally:
        with sin_triggers() as c:
            c.execute(text("UPDATE bitacora_auditoria SET usuario = :u WHERE id = :id"),
                      {"u": viejo.usuario, "id": viejo.id})
    assert _verificar()["integra"]


def test_registros_concurrentes_no_bifurcan_la_cadena():
    antes = _verificar()["ultima_secuencia"]
    errores: list[str] = []

    def trabajar():
        try:
            for _ in range(8):
                _registrar(1)
        except Exception as exc:  # noqa: BLE001
            errores.append(repr(exc))

    hilos = [threading.Thread(target=trabajar) for _ in range(4)]
    for h in hilos:
        h.start()
    for h in hilos:
        h.join(timeout=60)

    resultado = _verificar()
    assert not errores, errores
    assert resultado["integra"] and resultado["ultima_secuencia"] == antes + 32, resultado


# ── Ficha de custodia (PLUS-06, HU-41) ────────────────────────────────


def _expediente_con_documento(rfc: str) -> tuple[str, dict]:
    db = SessionLocal()
    try:
        titular = TitularExpediente(rfc=rfc, razon_social=f"Custodia {rfc} S.A. de C.V.", tipo_persona="MORAL")
        db.add(titular)
        db.commit()
        titular_id = titular.id
    finally:
        db.close()
    exp_id = client.post("/api/v1/expedientes", json={"titular_id": titular_id}).json()["id"]
    return exp_id, _subir_csf(exp_id, "csf_v1.pdf")


def _subir_csf(exp_id: str, nombre: str) -> dict:
    r = client.post("/api/v1/documentos/upload", data={"expediente_id": exp_id, "tipo_documento": "CSF_SAT"},
                    files={"file": (nombre, b"%PDF-1.4 " + nombre.encode() + os.urandom(8), "application/pdf")})
    assert r.status_code == 201, r.text
    return r.json()


def test_la_ficha_de_custodia_reune_el_ciclo_de_vida_del_documento():
    exp_id, v1 = _expediente_con_documento("CUS010101AAA")
    dictamen = client.post("/api/v1/dictamenes", json={
        "documento_id": v1["id"], "estatus": "EN_REVISION", "nivel_confianza": 0.8, "observaciones": ["Prueba."]}).json()
    assert client.post(f"/api/v1/dictamenes/{dictamen['id']}/ratificar", json={
        "nuevo_estatus": "CONFORME", "justificacion": "Constancia revisada contra el portal del SAT.",
        "usuario": "auditora.p3"}).status_code == 200
    _subir_csf(exp_id, "csf_v2.pdf")

    xlsx = client.get(f"/api/v1/reportes/documentos/{v1['id']}/custodia", params={"formato": "xlsx"})
    pdf = client.get(f"/api/v1/reportes/documentos/{v1['id']}/custodia", params={"formato": "pdf"})

    assert xlsx.status_code == 200 and pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")
    libro = load_workbook(BytesIO(xlsx.content))
    assert libro.sheetnames == ["Ficha", "Versiones", "Dictámenes", "Eventos", "Cadena"]
    ficha = {f[0]: f[1] for f in libro["Ficha"].iter_rows(min_row=3, values_only=True) if f[0]}
    assert ficha["Huella SHA-256 de la carga"] == v1["hash_sha256"] == ficha["Huella recalculada al emitir"]
    assert ficha["Integridad (RIE-03)"].startswith("Íntegro") and "1 de 2" in ficha["Versión (RF-04)"]
    assert "ÍNTEGRA" in ficha["Cadena de la bitácora"] and "no sustituye" in ficha["Alcance normativo"].lower()
    versiones = list(libro["Versiones"].iter_rows(min_row=2, values_only=True))
    assert [v[4] for v in versiones] == ["Sustituida", "Vigente"]
    [fila_dictamen] = list(libro["Dictámenes"].iter_rows(min_row=2, values_only=True))
    assert "auditora.p3" in fila_dictamen[5] and fila_dictamen[4] == "Coincide con la carga"
    eventos = list(libro["Eventos"].iter_rows(min_row=2, values_only=True))
    assert {"INGESTA_DOCUMENTO", "DICTAMEN_IA", "RATIFICACION_HUMANA"} <= {e[2] for e in eventos}
    assert all(isinstance(e[0], int) and len(e[5]) == 64 for e in eventos), "Cada evento con su sello"
    texto = "".join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages)
    assert "Ficha de Custodia Documental" in texto and "NOM-151" in texto and v1["hash_sha256"][:20] in texto

    exportaciones = client.get("/api/v1/bitacora", params={"documento_id": v1["id"],
                                                          "evento": "EXPORTACION_REPORTE"}).json()
    assert len(exportaciones) == 2 and all(e["detalle"]["reporte"] == "FICHA_CUSTODIA" and e["hash_evento"]
                                           for e in exportaciones)
    assert _verificar()["integra"]


def test_la_ficha_reporta_un_documento_alterado():
    _, doc = _expediente_con_documento("CUS020202AAA")
    with open(os.path.join(BOVEDA, f"{doc['hash_sha256'][:12]}_csf_v1.pdf"), "wb") as archivo:
        archivo.write(b"%PDF-1.4 contenido manipulado")

    r = client.get(f"/api/v1/reportes/documentos/{doc['id']}/custodia", params={"formato": "xlsx"})

    assert r.status_code == 200, r.text
    ficha = {f[0]: f[1] for f in load_workbook(BytesIO(r.content))["Ficha"].iter_rows(min_row=3, values_only=True) if f[0]}
    assert ficha["Integridad (RIE-03)"].startswith("ALTERADO")
    assert ficha["Huella recalculada al emitir"] != doc["hash_sha256"]
    assert client.get(f"/api/v1/documentos/{doc['id']}").json()["estado_integridad"] == "ALTERADO"


def test_ficha_de_un_documento_inexistente():
    assert client.get("/api/v1/reportes/documentos/no-existe/custodia").status_code == 404
