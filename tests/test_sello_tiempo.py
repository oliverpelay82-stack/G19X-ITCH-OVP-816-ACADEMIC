"""
P7 — Sello de tiempo de un tercero sobre la bitácora (PRD §21.7: PLUS-22, HU-58).

  1. Apagado por omisión: sin SELLO_TIEMPO_URL no sale ninguna llamada, el
     endpoint responde 409 y la reevaluación programada no sella.
  2. Con una autoridad RFC 3161 simulada: solo viaja la huella del último
     eslabón; el sello se guarda con su token, deja el evento SELLO_TIEMPO,
     la verificación de la cadena lo reporta y el token se descarga. Sin
     eventos nuevos no se pide otro.
  3. La tabla sellos_tiempo solo admite inserciones.
  4. Una autoridad caída, un sello rechazado o uno que no cubre la huella no
     guardan nada, y la reevaluación programada termina igual.
  5. Si la bitácora se reescribe después del sello, la verificación lo delata.
  6. El reporte global imprime el último sello.

Aislamiento: base de datos de pruebas plurione_compliance_test (tests/bd_pruebas.py).
"""
import os
import sys
from io import BytesIO

import pytest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bd_pruebas  # noqa: E402

BOVEDA = bd_pruebas.preparar_entorno()
os.environ["REEVALUACION_INTERVALO_HORAS"] = "0"
os.environ.pop("SELLO_TIEMPO_URL", None)

import requests  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import text  # noqa: E402
from sqlalchemy.exc import DBAPIError  # noqa: E402
from main import app  # noqa: E402
from clients import rules_client  # noqa: E402
from config.database import SessionLocal, engine, init_db  # noqa: E402
from models.bitacora import BitacoraAuditoria  # noqa: E402
from models.sello_tiempo import SelloTiempo  # noqa: E402
from motor import bitacora, reevaluacion_programada, sello_tiempo as st  # noqa: E402

bd_pruebas.verificar_aislamiento(engine)
init_db()
rules_client.obtener_documentos_obligatorios = lambda *_: rules_client.InventarioObligatorio(
    documentos={"CSF_SAT": {"codigo_regla": "RIE-01", "severidad": "CRITICA"}},
    fuente=rules_client.FUENTE_CATALOGO,
)
client = TestClient(app)

AUTORIDAD = "https://usuario:secreto@tsa.ficticia.test/tsr?cliente=7"
POLITICA = "1.3.6.1.4.1.99999.1.2"


def respuesta_simulada(solicitud: bytes, *, estado: int = 0, serie: int = 2**70 + 5,
                       fecha: str = "20261001120000.25Z", otra_huella: bool = False) -> bytes:
    """TimeStampResp de una autoridad ficticia que sella lo que se le pide (sin firma: la comprueba openssl)."""
    nodo, _ = st._leer(solicitud)
    _version, imprint, nonce, *_ = st._hijos(nodo.contenido)
    if otra_huella:
        imprint = st._secuencia(st._secuencia(st._oid(st.OID_SHA256), b"\x05\x00"), st._tlv(0x04, b"\x11" * 32))
    else:
        imprint = imprint.crudo
    tst = st._secuencia(st._entero(1), st._oid(POLITICA), imprint, st._entero(serie),
                        st._tlv(0x18, fecha.encode()), nonce.crudo)
    encapsulado = st._secuencia(st._oid(st.OID_TST_INFO), st._tlv(0xA0, st._tlv(0x04, tst)))
    firmado = st._secuencia(st._entero(3), st._tlv(0x31, st._secuencia(st._oid(st.OID_SHA256), b"\x05\x00")),
                            encapsulado, st._tlv(0x31, b""))
    token = st._secuencia(st._oid(st.OID_SIGNED_DATA), st._tlv(0xA0, firmado))
    return st._secuencia(st._secuencia(st._entero(estado)), token)


class _Respuesta:
    def __init__(self, contenido: bytes, codigo: int = 200):
        self.content, self.status_code = contenido, codigo


@pytest.fixture
def autoridad(monkeypatch):
    """Autoridad simulada; registra lo que el motor le envía."""
    llamadas: list[dict] = []
    ajustes: dict = {}

    def post(url, data=None, headers=None, timeout=None):
        llamadas.append({"url": url, "data": data, "headers": dict(headers or {})})
        if ajustes.get("caida"):
            raise requests.ConnectionError("sin ruta a la autoridad")
        return _Respuesta(respuesta_simulada(data, **ajustes.get("respuesta", {})))

    monkeypatch.setattr(st.requests, "post", post)
    monkeypatch.setenv("SELLO_TIEMPO_URL", AUTORIDAD)
    return llamadas, ajustes


def _registrar_evento() -> BitacoraAuditoria:
    db = SessionLocal()
    try:
        e = bitacora.registrar(db, "PRUEBA_SELLO_TIEMPO", usuario="pruebas.p7", detalle={"dato": "confidencial"})
        db.commit()
        db.refresh(e)
        return e
    finally:
        db.close()


def _contar(modelo, **filtros) -> int:
    db = SessionLocal()
    try:
        return db.query(modelo).filter_by(**filtros).count()
    finally:
        db.close()


def test_apagado_por_omision_no_llama_a_nadie(monkeypatch):
    llamadas = []
    monkeypatch.setattr(st.requests, "post", lambda *a, **k: llamadas.append(a) or _Respuesta(b""))
    _registrar_evento()

    r = client.post("/api/v1/bitacora/sello-tiempo")
    assert r.status_code == 409 and "SELLO_TIEMPO_URL" in r.json()["detail"]
    assert reevaluacion_programada.sellar_bitacora() is None
    assert llamadas == []


def test_sella_el_ultimo_eslabon_y_solo_envia_su_huella(autoridad):
    llamadas, _ = autoridad
    eslabon = _registrar_evento()
    sellos_antes = _contar(SelloTiempo)

    r = client.post("/api/v1/bitacora/sello-tiempo")
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["nuevo"] is True
    sello = cuerpo["sello"]
    assert sello["secuencia"] == eslabon.secuencia and sello["hash_evento"] == eslabon.hash_evento
    assert sello["numero_serie"] == str(2**70 + 5) and sello["politica"] == POLITICA
    assert sello["fecha_sello"] == "2026-10-01T12:00:00.250000Z" and sello["origen"] == st.ORIGEN_A_DEMANDA
    # La autoridad se imprime sin credenciales ni parámetros.
    assert sello["autoridad"] == "https://tsa.ficticia.test/tsr"
    assert _contar(SelloTiempo) == sellos_antes + 1

    # Solo viajó la solicitud RFC 3161 con la huella: ni el detalle del evento ni el folio interno.
    assert len(llamadas) == 1
    enviado = llamadas[0]["data"]
    assert bytes.fromhex(eslabon.hash_evento) in enviado and b"confidencial" not in enviado and len(enviado) < 80
    assert llamadas[0]["headers"] == {"Content-Type": "application/timestamp-query"}

    # Evento SELLO_TIEMPO sellado en la cadena, con lo que cubrió.
    evento = client.get("/api/v1/bitacora", params={"evento": "SELLO_TIEMPO", "limite": 1}).json()[0]
    assert evento["detalle"]["secuencia_sellada"] == eslabon.secuencia
    assert evento["detalle"]["sello_id"] == sello["id"] and evento["secuencia"] == eslabon.secuencia + 1

    cadena = client.get("/api/v1/bitacora/verificar-cadena").json()
    assert cadena["integra"], cadena["hallazgos"]
    assert cadena["sello_tiempo"]["id"] == sello["id"] and cadena["sello_tiempo"]["coincide"] is True

    # El token se descarga tal como lo emitió la autoridad.
    token = client.get(f"/api/v1/bitacora/sellos-tiempo/{sello['id']}/token")
    assert token.status_code == 200 and token.content[0] == 0x30
    assert bytes.fromhex(eslabon.hash_evento) in token.content
    listado = client.get("/api/v1/bitacora/sellos-tiempo").json()
    assert listado[0]["id"] == sello["id"] and listado[0]["coincide"] is True

    # Sin eventos nuevos (el último eslabón es el SELLO_TIEMPO) no se pide otro.
    otra = client.post("/api/v1/bitacora/sello-tiempo").json()
    assert otra["nuevo"] is False and otra["sello"]["id"] == sello["id"] and len(llamadas) == 1


def test_reevaluacion_programada_sella(autoridad):
    eslabon = _registrar_evento()
    resultado = reevaluacion_programada.sellar_bitacora()
    assert resultado.startswith("secuencia ")
    db = SessionLocal()
    try:
        ultimo = st.ultimo(db)
        assert ultimo.origen == st.ORIGEN_PROGRAMADO and ultimo.secuencia >= eslabon.secuencia
    finally:
        db.close()


def test_sellos_tiempo_solo_admite_inserciones(autoridad):
    _registrar_evento()
    client.post("/api/v1/bitacora/sello-tiempo")
    for sentencia in ("UPDATE sellos_tiempo SET numero_serie = '1'", "DELETE FROM sellos_tiempo"):
        with pytest.raises(DBAPIError, match="inmutable"):
            with engine.begin() as c:
                c.execute(text(sentencia))


@pytest.mark.parametrize("ajuste, texto", [
    ({"caida": True}, "no respondió"),
    ({"respuesta": {"estado": 2}}, "rechazado"),
    ({"respuesta": {"otra_huella": True}}, "no cubre la huella"),
])
def test_falla_de_la_autoridad_no_guarda_nada_ni_detiene_la_reevaluacion(autoridad, ajuste, texto):
    _, ajustes = autoridad
    ajustes.update(ajuste)
    _registrar_evento()
    sellos, eventos = _contar(SelloTiempo), _contar(BitacoraAuditoria, evento="SELLO_TIEMPO")

    r = client.post("/api/v1/bitacora/sello-tiempo")
    assert r.status_code == 502 and texto in r.json()["detail"]

    resultado = reevaluacion_programada.reevaluar_todos()
    assert resultado.sello.startswith("fallido") and texto in resultado.sello
    assert _contar(SelloTiempo) == sellos and _contar(BitacoraAuditoria, evento="SELLO_TIEMPO") == eventos
    assert client.get("/api/v1/bitacora/verificar-cadena").json()["integra"]


def test_reescritura_posterior_al_sello_se_detecta(autoridad):
    """Un sello sobre una huella que la cadena ya no tiene: la bitácora se reescribió después."""
    eslabon = _registrar_evento()
    ajeno = SelloTiempo(secuencia=eslabon.secuencia, evento_id=eslabon.id, hash_evento="ab" * 32,
                        autoridad="https://tsa.ficticia.test/tsr", numero_serie="1",
                        fecha_sello=eslabon.creado_en, token_der=b"\x30\x00", origen=st.ORIGEN_A_DEMANDA)
    db = SessionLocal()
    try:
        db.add(ajeno)
        db.commit()
        ajeno_id = ajeno.id
    finally:
        db.close()
    try:
        cadena = client.get("/api/v1/bitacora/verificar-cadena").json()
        assert not cadena["integra"]
        assert any(h["secuencia"] == eslabon.secuencia and "se reescribió" in h["problema"] for h in cadena["hallazgos"])
    finally:
        with engine.begin() as c:
            c.execute(text("ALTER TABLE sellos_tiempo DISABLE TRIGGER trg_sello_tiempo_inmutable"))
            c.execute(text("DELETE FROM sellos_tiempo WHERE id = :i"), {"i": ajeno_id})
            c.execute(text("ALTER TABLE sellos_tiempo ENABLE TRIGGER trg_sello_tiempo_inmutable"))
    assert client.get("/api/v1/bitacora/verificar-cadena").json()["integra"]


def test_reporte_global_imprime_el_ultimo_sello(autoridad):
    _registrar_evento()
    sello = client.post("/api/v1/bitacora/sello-tiempo").json()["sello"]
    r = client.get("/api/v1/reportes/cumplimiento", params={"formato": "xlsx"})
    assert r.status_code == 200
    celdas = [str(c.value) for fila in load_workbook(BytesIO(r.content))["Resumen"].iter_rows() for c in fila if c.value]
    assert "Sello de tiempo de un tercero" in celdas
    assert any(f"secuencia {sello['secuencia']} · coincide con la cadena" in c for c in celdas)
    assert any(sello["numero_serie"] in c for c in celdas)


def test_oid_y_longitudes_largas_ida_y_vuelta():
    for oid in (st.OID_SHA256, st.OID_TST_INFO, POLITICA, "1.2.840.113549.1.9.16.2.47"):
        assert st._valor_oid(st._leer(st._oid(oid))[0]) == oid
    largo = st._tlv(0x04, b"x" * 300)
    nodo, fin = st._leer(largo)
    assert nodo.contenido == b"x" * 300 and fin == len(largo)
    with pytest.raises(st.SelloTiempoError):
        st.interpretar_respuesta(b"\x30\x05\x02\x01", "00" * 32, 1)
