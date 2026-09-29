"""
Pruebas del Sprint 1-2 pendiente (core-service): tipo de expediente, anexos
obligatorios, vigencia recalculada, alertas con recomendación, búsqueda y
reportes exportables.

  1. Tipo de expediente (RF-01, HU-01): se guarda, y cada tipo exige su propio
     inventario del catálogo de reglas (RF-06); el respaldo DOC-02 también
     distingue el tipo.
  2. Anexos (RIE-02, RF-07, HU-10): el anexo solo se exige si el documento
     principal está presente; su ausencia deja el expediente NO_CONFORME.
  3. Vigencia recalculada (RN-02, RF-19): un documento auditado vigente que
     vence después pasa a NO_CONFORME; uno próximo a vencer, a EN_REVISION.
  4. Alertas con recomendación (RF-18, RF-21, HU-15, HU-16) y bandeja /alertas.
  5. Búsqueda y filtros de expedientes (RF-05, HU-04).
  6. Reportes PDF y Excel (RF-23, RF-24) y su registro en la bitácora.
  7. Verificación de contenido (RSC-01): las verificaciones del Nivel 4 quedan
     en reglas_aplicadas y cada elemento pendiente genera una alerta
     CONTENIDO_NO_LOCALIZADO con su recomendación hasta que un analista ratifica.
  8. Ejecución programada (DOC-02): la reevaluación diaria corrige un semáforo
     guardado desactualizado y deja el cambio en la bitácora.

Aislamiento: base plurione_compliance_test y bóveda temporal (bd_pruebas.py).

Uso:  python tests/test_tipos_anexos_alertas.py   (requiere PostgreSQL)
"""
import io
import os
import re
import sys
from datetime import date, timedelta
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bd_pruebas  # noqa: E402

BOVEDA = bd_pruebas.preparar_entorno()

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from main import app  # noqa: E402
from clients import rules_client  # noqa: E402
from config.database import SessionLocal, engine, init_db  # noqa: E402
from models.titular import TitularExpediente  # noqa: E402
from models.expediente import Expediente  # noqa: E402
from motor import cumplimiento, reevaluacion_programada  # noqa: E402

bd_pruebas.verificar_aislamiento(engine)
init_db()
client = TestClient(app)
_resultados: list[bool] = []

OBTENER_REAL = rules_client.obtener_documentos_obligatorios

# Catálogo simulado del rules-service por tipo de expediente (RF-06) con el
# anexo RIE-02 del contrato, igual que el seed.
CATALOGO = {
    "CORPORATIVO": ("CSF_SAT", "COMP_DOMICILIO", "ACTA_CONSTITUTIVA"),
    "JURIDICO": ("CONTRATO_SERVICIOS", "ACUERDO_CONFIDENCIALIDAD", "PODER_NOTARIAL"),
    "REGULATORIO": ("LICENCIA_OPERACION", "PERMISO_SECTORIAL", "POLIZA_SEGURO", "REGISTRO_SECTORIAL"),
}
consultas_catalogo: list[str] = []


def catalogo_por_tipo(tipo_expediente: str = "CORPORATIVO") -> rules_client.InventarioObligatorio:
    consultas_catalogo.append(tipo_expediente)
    return rules_client.InventarioObligatorio(
        documentos={t: {"codigo_regla": f"RIE-01-{t[:4]}", "severidad": "CRITICA"} for t in CATALOGO[tipo_expediente]},
        fuente=rules_client.FUENTE_CATALOGO,
        anexos=[{"tipo_documento": "ANEXO_CONTRATO", "documento_principal": "CONTRATO_SERVICIOS",
                 "codigo_regla": "RIE-02-CONT-ANX", "severidad": "ALTA"}],
    )


def verificar(condicion: bool, titulo: str, detalle: str = "") -> None:
    _resultados.append(bool(condicion))
    print(f"  {'[ OK ]' if condicion else '[FALLA]'} {titulo}" + (f"  -> {detalle}" if detalle else ""))


def titular(rfc: str, razon_social: str | None = None) -> str:
    db = SessionLocal()
    try:
        t = TitularExpediente(rfc=rfc, razon_social=razon_social or f"Titular {rfc} S.A. de C.V.", tipo_persona="MORAL")
        db.add(t)
        db.commit()
        return t.id
    finally:
        db.close()


def abrir(titular_id: str, tipo: str | None = None, nombre: str | None = None) -> dict:
    cuerpo = {"titular_id": titular_id}
    if tipo:
        cuerpo["tipo_expediente"] = tipo
    if nombre:
        cuerpo["nombre"] = nombre
    r = client.post("/api/v1/expedientes", json=cuerpo)
    assert r.status_code == 201, r.text
    return r.json()


def subir(exp_id: str, tipo: str, nombre: str) -> dict:
    r = client.post(
        "/api/v1/documentos/upload",
        data={"expediente_id": exp_id, "tipo_documento": tipo},
        files={"file": (nombre, b"%PDF-1.4 " + nombre.encode() + os.urandom(8), "application/pdf")},
    )
    assert r.status_code == 201, r.text
    return r.json()


def dictaminar(doc_id: str, estatus: str = "CONFORME", **extra) -> dict:
    cuerpo = {"documento_id": doc_id, "estatus": estatus, "nivel_confianza": 0.95,
              "observaciones": ["Prueba automatizada."], **extra}
    r = client.post("/api/v1/dictamenes", json=cuerpo)
    assert r.status_code == 201, r.text
    return r.json()


def expediente(exp_id: str) -> dict:
    return client.get(f"/api/v1/expedientes/{exp_id}").json()


def codigos(datos: dict) -> set[str]:
    return {a["codigo"] for a in datos["evaluacion"]["alertas"]}


def prueba_tipo_expediente() -> dict:
    print("\n[1] Tipo de expediente e inventario por tipo (RF-01, RF-06)")
    t = titular("TIP010101AAA", "Constructora del Norte S.A. de C.V.")
    corporativo = abrir(t)
    verificar(corporativo["tipo_expediente"] == "CORPORATIVO", "Sin tipo explícito, el expediente es CORPORATIVO",
              corporativo["tipo_expediente"])
    juridico = abrir(t, "JURIDICO", "Contrato de mantenimiento 2026")
    verificar(juridico["tipo_expediente"] == "JURIDICO" and juridico["nombre"] == "Contrato de mantenimiento 2026",
              "El tipo y el nombre descriptivo se guardan")
    verificar(juridico["codigo_expediente"] != corporativo["codigo_expediente"],
              "Dos expedientes del mismo titular en el mismo minuto reciben folios distintos")
    verificar(juridico["creado_en"] is not None, "La fecha de apertura se registra (RF-05)")
    exigidos = {o["tipo_documento"] for o in juridico["evaluacion"]["obligatorios"]}
    verificar(exigidos == set(CATALOGO["JURIDICO"]), "El jurídico exige contrato, NDA y poder notarial",
              ", ".join(sorted(exigidos)))
    verificar("CSF_SAT" not in exigidos, "El jurídico no exige la CSF del corporativo")
    regulatorio = abrir(t, "REGULATORIO")
    verificar(len(regulatorio["evaluacion"]["obligatorios"]) == 4, "El regulatorio exige su propio inventario (4)")
    invalido = client.post("/api/v1/expedientes", json={"titular_id": t, "tipo_expediente": "FISCAL"})
    verificar(invalido.status_code == 422, "Un tipo fuera del catálogo se rechaza (422)", f"HTTP {invalido.status_code}")

    rules_client.obtener_documentos_obligatorios = OBTENER_REAL
    url_real = rules_client.RULES_BASE_URL
    rules_client.RULES_BASE_URL = "http://127.0.0.1:9"
    try:
        respaldo = rules_client.obtener_documentos_obligatorios("JURIDICO")
    finally:
        rules_client.RULES_BASE_URL = url_real
        rules_client.obtener_documentos_obligatorios = catalogo_por_tipo
    verificar(respaldo.fuente == rules_client.FUENTE_RESPALDO and "CONTRATO_SERVICIOS" in respaldo.documentos
              and respaldo.anexos, "El respaldo DOC-02 también distingue el tipo y trae los anexos",
              ", ".join(respaldo.documentos))
    return {"titular": t, "juridico": juridico, "corporativo": corporativo, "regulatorio": regulatorio}


def prueba_anexos(juridico: dict) -> None:
    print("\n[2] Anexos obligatorios vinculados (RIE-02, RF-07, HU-10)")
    exp_id = juridico["id"]
    nda = subir(exp_id, "ACUERDO_CONFIDENCIALIDAD", "nda.pdf")
    poder = subir(exp_id, "PODER_NOTARIAL", "poder.pdf")
    datos = expediente(exp_id)
    verificar(all(o["documento_principal"] is None for o in datos["evaluacion"]["obligatorios"]),
              "Sin contrato, el anexo del contrato no se exige")

    contrato = subir(exp_id, "CONTRATO_SERVICIOS", "contrato.pdf")
    for doc in (nda, poder, contrato):
        dictaminar(doc["id"])
    datos = expediente(exp_id)
    anexos = [o for o in datos["evaluacion"]["obligatorios"] if o["documento_principal"]]
    verificar(len(anexos) == 1 and anexos[0]["tipo_documento"] == "ANEXO_CONTRATO" and not anexos[0]["presente"],
              "Con el contrato presente, su anexo legal pasa a ser obligatorio")
    verificar(datos["estado_cumplimiento"] == "NO_CONFORME",
              "Contrato sin su anexo: el expediente queda NO_CONFORME (RIE-02)", datos["estado_cumplimiento"])
    verificar(datos["evaluacion"]["anexos_faltantes"][0]["documento_principal"] == "CONTRATO_SERVICIOS",
              "La evaluación indica qué anexo falta y de qué documento")
    verificar(datos["porcentaje_completitud"] == 75.0, "La completitud cuenta el anexo exigido (3 de 4)",
              f"{datos['porcentaje_completitud']}%")
    alerta = next((a for a in datos["evaluacion"]["alertas"] if a["codigo"] == "ANEXO_FALTANTE"), None)
    verificar(alerta is not None and "Contrato" in alerta["recomendacion"] and alerta["regla"] == "RIE-02-CONT-ANX",
              "Alerta ANEXO_FALTANTE con recomendación específica", alerta["recomendacion"] if alerta else "sin alerta")

    anexo = subir(exp_id, "ANEXO_CONTRATO", "anexo_a.pdf")
    dictaminar(anexo["id"])
    datos = expediente(exp_id)
    fila = next(d for d in datos["evaluacion"]["documentos"] if d["tipo_documento"] == "ANEXO_CONTRATO")
    verificar(datos["estado_cumplimiento"] == "CONFORME" and fila["anexo_de"] == "CONTRATO_SERVICIOS",
              "Con el anexo auditado conforme, el expediente pasa a CONFORME", datos["estado_cumplimiento"])
    verificar(not datos["evaluacion"]["alertas"], "Un expediente conforme no tiene alertas")


def prueba_vigencia_y_alertas(ctx: dict) -> None:
    print("\n[3] Vigencia recalculada contra la fecha actual (RN-02, RF-19)")
    exp_id = ctx["corporativo"]["id"]
    hoy = date.today()
    csf = subir(exp_id, "CSF_SAT", "csf.pdf")
    dom = subir(exp_id, "COMP_DOMICILIO", "recibo.pdf")
    acta = subir(exp_id, "ACTA_CONSTITUTIVA", "acta.pdf")
    vigencia_rcv01 = {"disponible": True, "resultado": "VIGENTE", "evaluaciones": [
        {"codigo_regla": "RCV-01", "resultado": "VIGENTE", "limite_dias": 30, "criterio": "DIAS_PARA_VENCIMIENTO",
         "dias_calculados": 200, "detalle": "Vigente."}]}
    dictaminar(csf["id"], vigencia=vigencia_rcv01, fecha_vencimiento=(hoy + timedelta(days=200)).isoformat())
    # Recibo emitido hace 100 días, auditado cuando aún estaba dentro de los 90 (sin criterio: dictamen previo).
    vigencia_rcv02 = {"disponible": True, "resultado": "VIGENTE", "evaluaciones": [
        {"codigo_regla": "RCV-02", "resultado": "VIGENTE", "limite_dias": 90, "dias_calculados": 60,
         "detalle": "60 días de antigüedad."}]}
    dictaminar(dom["id"], vigencia=vigencia_rcv02, fecha_emision=(hoy - timedelta(days=100)).isoformat())
    dictaminar(acta["id"])
    datos = expediente(exp_id)
    fila_dom = next(d for d in datos["evaluacion"]["documentos"] if d["tipo_documento"] == "COMP_DOMICILIO")
    verificar(fila_dom["vigencia"] == "VENCIDO" and fila_dom["estado"] == "NO_CONFORME",
              "Recibo auditado vigente que ya superó los 90 días: VENCIDO y NO_CONFORME", fila_dom["vigencia"])
    verificar(datos["estado_cumplimiento"] == "NO_CONFORME", "El expediente pasa a NO_CONFORME sin reauditar (RN-02)")
    alerta = next((a for a in datos["evaluacion"]["alertas"] if a["codigo"] == "DOCUMENTO_VENCIDO"), None)
    verificar(alerta is not None and "90 días" in alerta["recomendacion"],
              "Alerta DOCUMENTO_VENCIDO recomienda un comprobante no mayor a 90 días",
              alerta["recomendacion"] if alerta else "sin alerta")
    verificar(alerta is not None and alerta["regla"] == "RCV-02", "La alerta cita la regla que venció el recibo",
              alerta["regla"] if alerta else "sin alerta")
    # La regla citada es la del dictamen, no una deducida del tipo: una constancia
    # vencida por RCV-03 no puede quedar sin regla (antes solo se citaba RCV-02).
    dictamen_csf = SimpleNamespace(
        fecha_vencimiento_extraida=None, fecha_emision_extraida=hoy - timedelta(days=120),
        reglas_aplicadas=[{"codigo_regla": "RCV-03", "limite_dias": 90, "criterio": "ANTIGUEDAD_DESDE_EMISION"}],
    )
    vig_csf = cumplimiento.vigencia_actual(dictamen_csf, hoy)
    verificar(vig_csf.estado == "VENCIDO" and vig_csf.regla == "RCV-03",
              "Una constancia vencida por antigüedad cita RCV-03", f"{vig_csf.estado} {vig_csf.regla}")

    dom2 = subir(exp_id, "COMP_DOMICILIO", "recibo_nuevo.pdf")
    dictaminar(dom2["id"], vigencia=vigencia_rcv02, fecha_emision=(hoy - timedelta(days=10)).isoformat())
    datos = expediente(exp_id)
    verificar(datos["estado_cumplimiento"] == "CONFORME", "Recargar un recibo reciente recupera el expediente",
              datos["estado_cumplimiento"])

    csf2 = subir(exp_id, "CSF_SAT", "csf_por_vencer.pdf")
    dictaminar(csf2["id"], vigencia=vigencia_rcv01, fecha_vencimiento=(hoy + timedelta(days=4)).isoformat())
    datos = expediente(exp_id)
    fila_csf = next(d for d in datos["evaluacion"]["documentos"] if d["tipo_documento"] == "CSF_SAT")
    verificar(fila_csf["vigencia"] == "PROXIMO_A_VENCER" and fila_csf["dias_para_vencer"] == 4,
              "CSF que vence en 4 días: PROXIMO_A_VENCER", fila_csf["vigencia"])
    verificar(datos["estado_cumplimiento"] == "EN_REVISION", "Documento próximo a vencer: semáforo amarillo (RF-19)",
              datos["estado_cumplimiento"])
    alerta = next((a for a in datos["evaluacion"]["alertas"] if a["codigo"] == "PROXIMO_A_VENCER"), None)
    verificar(alerta is not None and alerta["severidad"] == "ALTA",
              "A 4 días del vencimiento la alerta es ALTA (umbral de 5 días)", alerta["severidad"] if alerta else "-")
    verificar(cumplimiento._severidad_por_anticipacion(20) == "MEDIA"
              and cumplimiento._severidad_por_anticipacion(10) == "MEDIA_ALTA"
              and cumplimiento._severidad_por_anticipacion(45) is None,
              "Umbrales de anticipación 30 / 15 / 5 días (RF-17)")

    print("\n[4] Alertas con recomendación y bandeja general (RF-18, RF-21)")
    regulatorio = expediente(ctx["regulatorio"]["id"])
    faltantes = [a for a in regulatorio["evaluacion"]["alertas"] if a["codigo"] == "DOCUMENTO_FALTANTE"]
    verificar(len(faltantes) == 4 and all(a["recomendacion"] for a in faltantes),
              "Cada documento faltante genera una alerta con recomendación", f"{len(faltantes)} alertas")
    licencia = subir(ctx["regulatorio"]["id"], "LICENCIA_OPERACION", "licencia.pdf")
    datos = expediente(ctx["regulatorio"]["id"])
    verificar("PENDIENTE_AUDITORIA" in codigos(datos), "Un obligatorio cargado sin dictamen genera PENDIENTE_AUDITORIA")
    dictaminar(licencia["id"], "EN_REVISION", nivel_confianza=0.75)
    datos = expediente(ctx["regulatorio"]["id"])
    verificar("PENDIENTE_RATIFICACION" in codigos(datos), "Un dictamen en revisión genera PENDIENTE_RATIFICACION")

    bandeja = client.get("/api/v1/alertas").json()
    verificar(bandeja and {"expediente_id", "codigo_expediente", "razon_social", "recomendacion"} <= set(bandeja[0]),
              "GET /alertas consolida las alertas con su expediente", f"{len(bandeja)} alertas")
    orden = ["CRITICA", "ALTA", "MEDIA_ALTA", "MEDIA"]
    verificar([orden.index(a["severidad"]) for a in bandeja] == sorted(orden.index(a["severidad"]) for a in bandeja),
              "La bandeja se ordena de la más grave a la menos grave")
    criticas = client.get("/api/v1/alertas", params={"severidad": "CRITICA"}).json()
    verificar(criticas and all(a["severidad"] == "CRITICA" for a in criticas), "Filtro por severidad")
    solo_reg = client.get("/api/v1/alertas", params={"tipo_expediente": "REGULATORIO"}).json()
    verificar(solo_reg and all(a["tipo_expediente"] == "REGULATORIO" for a in solo_reg), "Filtro por tipo de expediente")


def prueba_busqueda(ctx: dict) -> None:
    print("\n[5] Búsqueda y filtros de expedientes (RF-05, HU-04)")
    listar = lambda **p: client.get("/api/v1/expedientes", params=p).json()  # noqa: E731
    verificar(len(listar(q="mantenimiento")) == 1, "Búsqueda por nombre del expediente")
    verificar(len(listar(q="constructora del norte")) == 3, "Búsqueda por razón social (sin distinguir mayúsculas)")
    verificar(len(listar(q="TIP010101")) == 3, "Búsqueda por RFC")
    verificar(all(e["tipo_expediente"] == "JURIDICO" for e in listar(tipo_expediente="JURIDICO")), "Filtro por tipo")
    hoy = date.today().isoformat()
    verificar(len(listar(desde=hoy, hasta=hoy)) >= 3, "Filtro por fecha de apertura (hoy)")
    verificar(listar(desde=(date.today() + timedelta(days=1)).isoformat()) == [], "Sin aperturas futuras")
    verificar(all(e["estado_cumplimiento"] == "CONFORME" for e in listar(estado="CONFORME")), "Filtro por estatus")
    verificar(client.get("/api/v1/expedientes", params={"tipo_expediente": "X"}).status_code == 422,
              "Un tipo inválido en el filtro responde 422")


def prueba_reportes(ctx: dict) -> None:
    print("\n[6] Reportes exportables (RF-23, RF-24)")
    exp_id = ctx["juridico"]["id"]
    # La base de pruebas es compartida: se cuentan solo las exportaciones de esta sección.
    previas = {e["id"] for e in client.get("/api/v1/bitacora", params={"evento": "EXPORTACION_REPORTE",
                                                                      "limite": 500}).json()}
    pdf = client.get(f"/api/v1/reportes/expedientes/{exp_id}", params={"formato": "pdf"})
    verificar(pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")
              and pdf.headers["content-type"] == "application/pdf"
              and "attachment" in pdf.headers.get("content-disposition", ""),
              "Dictamen del expediente en PDF", f"{len(pdf.content)} bytes")
    # RNF-08: el dictamen lo leen personas sin conocimientos técnicos. Los códigos
    # de regla van en su columna; en el texto no hay códigos internos ni referencias.
    texto_pdf = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    jerga = re.findall(r"\b[A-Z]{2,}_[A-Z_]+\b|\((?:RIE|RCV|RSC|RN|RF)-[^)]*\)|Nivel(?:es)? [1-5]", texto_pdf)
    verificar(not jerga, "El dictamen en PDF está en lenguaje llano", ", ".join(sorted(set(jerga))) or "sin códigos")
    xlsx = client.get(f"/api/v1/reportes/expedientes/{exp_id}", params={"formato": "xlsx"})
    libro = load_workbook(io.BytesIO(xlsx.content))
    verificar(xlsx.status_code == 200
              and libro.sheetnames == ["Dictamen", "Documentos", "Reglas", "Evidencia", "Alertas", "Bitácora"],
              "Dictamen del expediente en Excel con sus seis hojas", ", ".join(libro.sheetnames))
    # C6 (PRD §10.5): la hoja «Reglas» es la evidencia del índice de cumplimiento del DOC-05.
    indice = client.get(f"/api/v1/expedientes/{exp_id}").json()["evaluacion"]["indice_cumplimiento"]
    reglas = list(libro["Reglas"].iter_rows(min_row=2, values_only=True))
    ficha = {f[0]: f[1] for f in libro["Dictamen"].iter_rows(min_row=3, values_only=True) if f[0]}
    verificar(len(reglas) == indice["aplicables"]
              and sum(f[3] == "Cumplida" for f in reglas) == indice["cumplidas"]
              and ficha["Índice de cumplimiento"].startswith(f"{indice['porcentaje']:.1f} %"),
              "La hoja «Reglas» y la ficha cuadran con el índice de cumplimiento",
              f"{indice['cumplidas']} de {indice['aplicables']}: {ficha['Índice de cumplimiento']}")
    evidencia = list(libro["Evidencia"].iter_rows(min_row=2, values_only=True))
    verificar(len(evidencia) == 4 and all(len(fila[7]) == 64 for fila in evidencia),
              "La evidencia incluye cada dictamen con su huella SHA-256", f"{len(evidencia)} dictámenes")

    global_xlsx = client.get("/api/v1/reportes/cumplimiento", params={"formato": "xlsx"})
    libro = load_workbook(io.BytesIO(global_xlsx.content))
    filas = list(libro["Expedientes"].iter_rows(min_row=2, values_only=True))
    # La base de pruebas puede traer expedientes de otras suites (pytest tests/): se compara contra el total.
    total = len(client.get("/api/v1/expedientes").json())
    verificar(global_xlsx.status_code == 200 and len(filas) == total and total >= 3,
              "Reporte global en Excel con todos los expedientes", f"{len(filas)} filas")
    global_pdf = client.get("/api/v1/reportes/cumplimiento", params={"formato": "pdf"})
    verificar(global_pdf.status_code == 200 and global_pdf.content.startswith(b"%PDF-"), "Reporte global en PDF")
    texto_global = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(global_pdf.content)).pages)
    verificar("Expedientes conformes" in texto_global and "Índice de cumplimiento promedio" in texto_global,
              "El reporte global distingue expedientes conformes e índice de cumplimiento promedio")
    verificar(client.get("/api/v1/reportes/cumplimiento", params={"formato": "docx"}).status_code == 422,
              "Un formato no admitido responde 422")
    verificar(client.get("/api/v1/reportes/expedientes/no-existe").status_code == 404, "Expediente inexistente: 404")
    eventos = [e for e in client.get("/api/v1/bitacora", params={"evento": "EXPORTACION_REPORTE", "limite": 500}).json()
               if e["id"] not in previas]
    verificar(len(eventos) == 4 and {e["detalle"]["formato"] for e in eventos} == {"pdf", "xlsx"},
              "Cada exportación queda en la bitácora", f"{len(eventos)} eventos")


def prueba_verificacion_contenido() -> None:
    print("\n[7] Verificación de contenido del Nivel 4 (RSC-01)")
    exp_id = abrir(titular("VCO150310QK4", "Inmobiliaria Costa Grande S.A. de C.V."))["id"]
    documentos = {t: subir(exp_id, t, f"{t.lower()}.pdf") for t in CATALOGO["CORPORATIVO"]}
    for tipo in ("CSF_SAT", "COMP_DOMICILIO"):
        dictaminar(documentos[tipo]["id"])
    verificaciones = [
        {"codigo_regla": "RSC-ACTA-ESCR", "nombre_regla": "Escritura pública y notario", "familia_regla": "RSC",
         "severidad": "MEDIA_ALTA", "resultado": "CUMPLE", "cita": "ESCRITURA PUBLICA NUMERO 18,442",
         "observacion": "Escritura identificada.", "recomendacion": "Solicitar la copia certificada."},
        {"codigo_regla": "RSC-ACTA-RPC", "nombre_regla": "Inscripción en el Registro Público de Comercio",
         "familia_regla": "RSC", "severidad": "MEDIA_ALTA", "resultado": "NO_LOCALIZADO", "cita": None,
         "observacion": "No se menciona el folio mercantil.",
         "recomendacion": "Solicitar la boleta de inscripción en el Registro Público de Comercio."},
    ]
    dictamen = dictaminar(documentos["ACTA_CONSTITUTIVA"]["id"], "EN_REVISION", verificaciones=verificaciones,
                          requiere_revision_humana=True)
    reglas = {r["codigo_regla"]: r for r in dictamen["reglas_aplicadas"]}
    verificar(reglas.get("RSC-ACTA-ESCR", {}).get("cita") == "ESCRITURA PUBLICA NUMERO 18,442"
              and reglas.get("RSC-ACTA-RPC", {}).get("resultado") == "NO_LOCALIZADO",
              "Las verificaciones se guardan en reglas_aplicadas con su cita y resultado")
    datos = expediente(exp_id)
    contenido = [a for a in datos["evaluacion"]["alertas"] if a["codigo"] == cumplimiento.ALERTA_CONTENIDO]
    verificar(len(contenido) == 1 and contenido[0]["regla"] == "RSC-ACTA-RPC"
              and contenido[0]["recomendacion"].startswith("Solicitar la boleta"),
              "Solo el elemento pendiente genera CONTENIDO_NO_LOCALIZADO, con la recomendación del catálogo",
              f"{len(contenido)} alertas")
    ratificacion = next((a for a in datos["evaluacion"]["alertas"] if a["codigo"] == "PENDIENTE_RATIFICACION"), None)
    verificar(datos["estado_cumplimiento"] == "EN_REVISION" and ratificacion is not None,
              "Una omisión deja el expediente EN_REVISION, no NO_CONFORME", datos["estado_cumplimiento"])
    r = client.post(f"/api/v1/dictamenes/{dictamen['id']}/ratificar",
                    json={"nuevo_estatus": "CONFORME", "usuario": "analista.pruebas",
                          "justificacion": "El folio mercantil consta en la boleta anexa al acta, revisada en físico."})
    datos = expediente(exp_id)
    verificar(r.status_code == 200 and cumplimiento.ALERTA_CONTENIDO not in codigos(datos)
              and datos["estado_cumplimiento"] == "CONFORME",
              "Tras la ratificación humana la alerta desaparece (RN-05)", datos["estado_cumplimiento"])
    return exp_id


def prueba_reevaluacion_programada(exp_id: str) -> None:
    print("\n[8] Reevaluación programada del motor (ejecución programada, DOC-02)")
    # Semáforo guardado desactualizado, como el de un documento que venció sin que nadie consultara el expediente.
    db = SessionLocal()
    try:
        db.query(Expediente).filter(Expediente.id == exp_id).update({"estado_cumplimiento": "NO_CONFORME"})
        db.commit()
    finally:
        db.close()
    resultado = reevaluacion_programada.reevaluar_todos()
    # Con pytest tests/ la base trae expedientes de otras suites: se exige al menos el cambio provocado.
    verificar(resultado.evaluados >= 4 and resultado.cambios >= 1,
              "Reevalúa todos los expedientes y corrige el desactualizado",
              f"{resultado.evaluados} evaluados, {resultado.cambios} cambios")
    eventos = client.get("/api/v1/bitacora", params={"expediente_id": exp_id, "evento": "EVALUACION_REGLAS"}).json()
    ultimo = eventos[0]["detalle"] if eventos else {}
    verificar(ultimo.get("causa") == reevaluacion_programada.CAUSA and ultimo.get("estado_nuevo") == "CONFORME",
              "El cambio queda en la bitácora con la causa REEVALUACION_PROGRAMADA", ultimo.get("causa"))
    verificar(reevaluacion_programada.reevaluar_todos().cambios == 0, "Una segunda ejecución no registra cambios")
    os.environ["REEVALUACION_INTERVALO_HORAS"] = "0"
    desactivada = reevaluacion_programada.intervalo_horas() == 0
    os.environ["REEVALUACION_INTERVALO_HORAS"] = "abc"
    invalida = reevaluacion_programada.intervalo_horas() == 0
    os.environ.pop("REEVALUACION_INTERVALO_HORAS")
    verificar(desactivada and invalida and reevaluacion_programada.intervalo_horas() == 24,
              "Intervalo configurable: 24 h por defecto, 0 o un valor inválido la desactivan")


def run_tests() -> int:
    print("=== TIPO DE EXPEDIENTE, ANEXOS, VIGENCIA, ALERTAS, BÚSQUEDA Y REPORTES ===")
    rules_client.obtener_documentos_obligatorios = catalogo_por_tipo
    try:
        ctx = prueba_tipo_expediente()
        prueba_anexos(ctx["juridico"])
        prueba_vigencia_y_alertas(ctx)
        prueba_busqueda(ctx)
        prueba_reportes(ctx)
        prueba_reevaluacion_programada(prueba_verificacion_contenido())
    finally:
        rules_client.obtener_documentos_obligatorios = OBTENER_REAL
        bd_pruebas.limpiar_boveda(BOVEDA)

    aprobadas = sum(_resultados)
    print(f"\nResultado: {aprobadas}/{len(_resultados)} verificaciones aprobadas")
    return 0 if aprobadas == len(_resultados) else 1


def test_tipos_anexos_alertas():
    assert run_tests() == 0, "Hay verificaciones en [FALLA] (ver salida con -s)"


if __name__ == "__main__":
    sys.exit(run_tests())
