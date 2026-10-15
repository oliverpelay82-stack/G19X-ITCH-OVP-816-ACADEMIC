"""
PLUS-23 — Sello de integridad de cada reporte (ESTADO_DEL_DESARROLLO §6.9).

  1. El sello es el SHA-256 de los datos del reporte en forma canónica: los
     mismos datos dan el mismo sello, sin importar el orden de las claves, y
     cualquier cambio en los datos lo cambia.
  2. El dictamen del expediente, el reporte global y la ficha de custodia, en
     PDF y en Excel, llevan impreso el «Sello de integridad forense (SHA-256)».
  3. El evento EXPORTACION_REPORTE guarda ese sello y la huella del archivo
     emitido, con la que quien lo reciba comprueba que no cambió.

Aislamiento: base de datos de pruebas plurione_compliance_test (tests/bd_pruebas.py).
"""
import hashlib
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bd_pruebas  # noqa: E402

bd_pruebas.preparar_entorno()
os.environ["REEVALUACION_INTERVALO_HORAS"] = "0"

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from main import app  # noqa: E402
from clients import rules_client  # noqa: E402
from config.database import SessionLocal, engine, init_db  # noqa: E402
from models.titular import TitularExpediente  # noqa: E402
from motor import reportes  # noqa: E402

bd_pruebas.verificar_aislamiento(engine)
init_db()
rules_client.obtener_documentos_obligatorios = lambda *_: rules_client.InventarioObligatorio(
    documentos={"CSF_SAT": {"codigo_regla": "RIE-01", "severidad": "CRITICA"}},
    fuente=rules_client.FUENTE_CATALOGO,
)
auth = bd_pruebas.Autenticacion()
client = auth.cliente(app)


# ── El sello de los datos ────────────────────────────────────────────


@dataclass
class _Datos:
    folio: str
    emitido: datetime
    motivos: dict


def test_el_sello_depende_solo_de_los_datos():
    base = _Datos("EXP-00001", datetime(2026, 10, 3, 12, 0), {"a": 1, "b": [1, 2]})
    mismo = _Datos("EXP-00001", datetime(2026, 10, 3, 12, 0), {"b": [1, 2], "a": 1})
    otro = _Datos("EXP-00001", datetime(2026, 10, 3, 12, 0), {"a": 1, "b": [1, 3]})

    sello = reportes.sello_contenido(base)

    assert re.fullmatch(r"[0-9a-f]{64}", sello)
    assert sello == reportes.sello_contenido(mismo), "El orden de las claves no cambia el sello"
    assert sello != reportes.sello_contenido(otro)
    assert sello != reportes.sello_contenido(base, {"integra": True}), "La cadena también entra al sello"


# ── El sello en los reportes y en la bitácora ────────────────────────


def _documento(rfc: str) -> dict:
    db = SessionLocal()
    try:
        titular = TitularExpediente(rfc=rfc, razon_social=f"Sello {rfc} S.A. de C.V.", tipo_persona="MORAL")
        db.add(titular)
        db.commit()
        titular_id = titular.id
    finally:
        db.close()
    exp_id = client.post("/api/v1/expedientes", json={"titular_id": titular_id}).json()["id"]
    r = client.post("/api/v1/documentos/upload", data={"expediente_id": exp_id, "tipo_documento": "CSF_SAT"},
                    files={"file": ("csf.pdf", b"%PDF-1.4 sello " + os.urandom(8), "application/pdf")})
    assert r.status_code == 201, r.text
    return r.json()


def _sello_excel(contenido: bytes) -> str:
    libro = load_workbook(BytesIO(contenido))
    hoja = libro[libro.sheetnames[0]]
    [sello] = [fila[1] for fila in hoja.iter_rows(values_only=True) if fila[0] == reportes.ROTULO_SELLO]
    return sello


def _sellos_pdf(contenido: bytes) -> list[str]:
    paginas = PdfReader(BytesIO(contenido)).pages
    return [re.search(rf"{re.escape(reportes.ROTULO_SELLO)}: ([0-9a-f]{{64}})", p.extract_text()).group(1)
            for p in paginas]


def _ultima_exportacion(**filtro) -> dict:
    eventos = client.get("/api/v1/bitacora", params={"evento": "EXPORTACION_REPORTE", **filtro}).json()
    return max(eventos, key=lambda e: e["secuencia"])


@pytest.mark.parametrize("reporte", ["expediente", "global", "custodia"])
def test_cada_reporte_lleva_su_sello_y_la_bitacora_lo_respalda(reporte):
    doc = _documento({"expediente": "SEL010101AAA", "global": "SEL020202AAA", "custodia": "SEL030303AAA"}[reporte])
    ruta, filtro = {
        "expediente": (f"/api/v1/reportes/expedientes/{doc['expediente_id']}", {"expediente_id": doc["expediente_id"]}),
        "global": ("/api/v1/reportes/cumplimiento", {}),
        "custodia": (f"/api/v1/reportes/documentos/{doc['id']}/custodia", {"documento_id": doc["id"]}),
    }[reporte]

    for formato in ("xlsx", "pdf"):
        r = client.get(ruta, params={"formato": formato})
        assert r.status_code == 200, r.text

        if formato == "xlsx":
            sello = _sello_excel(r.content)
        else:
            sellos = _sellos_pdf(r.content)
            assert len(set(sellos)) == 1, "El mismo sello en el pie de cada página"
            sello = sellos[0]
        evento = _ultima_exportacion(**filtro)["detalle"]
        assert evento["formato"] == formato
        assert evento["sello_contenido"] == sello
        assert evento["hash_archivo"] == hashlib.sha256(r.content).hexdigest()
