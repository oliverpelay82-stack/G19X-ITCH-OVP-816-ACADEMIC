"""
P5 — Integridad permanente (ESTADO_DEL_DESARROLLO §6.5: PLUS-14, HU-44; PLUS-15, HU-45).

  1. El barrido programado detecta un binario alterado que nadie leyó: el
     documento queda ALTERADO con el mismo dictamen RIE-03 y el mismo evento
     ALERTA_INTEGRIDAD que una lectura, con la nota del barrido; los documentos
     íntegros no cambian.
  2. Un binario eliminado de la bóveda también se detecta.
  3. Cada barrido deja exactamente una constancia VERIFICACION_INTEGRIDAD
     sellada, y un documento ya ALTERADO no se vuelve a registrar.
  4. La reevaluación programada barre antes de reevaluar, y
     BARRIDO_INTEGRIDAD=false lo desactiva.
  5. El reporte del expediente y el global, en PDF y en Excel, imprimen el
     último sello de la cadena antes de su exportación (el hash_anterior del
     evento que registra esa exportación).

Aislamiento: base de datos de pruebas plurione_compliance_test (tests/bd_pruebas.py).
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import bd_pruebas  # noqa: E402

BOVEDA = bd_pruebas.preparar_entorno()
os.environ["REEVALUACION_INTERVALO_HORAS"] = "0"

from io import BytesIO  # noqa: E402

from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from main import app  # noqa: E402
from clients import rules_client  # noqa: E402
from config.database import SessionLocal, engine, init_db  # noqa: E402
from models.bitacora import BitacoraAuditoria  # noqa: E402
from models.dictamen import ORIGEN_RIE_03, DictamenAuditoria  # noqa: E402
from models.titular import TitularExpediente  # noqa: E402
from motor import integridad, reevaluacion_programada  # noqa: E402

bd_pruebas.verificar_aislamiento(engine)
init_db()
rules_client.obtener_documentos_obligatorios = lambda *_: rules_client.InventarioObligatorio(
    documentos={"CSF_SAT": {"codigo_regla": "RIE-01", "severidad": "CRITICA"}},
    fuente=rules_client.FUENTE_CATALOGO,
)
auth = bd_pruebas.Autenticacion()
client = auth.cliente(app)


def _documento(rfc: str) -> dict:
    db = SessionLocal()
    try:
        titular = TitularExpediente(rfc=rfc, razon_social=f"Barrido {rfc} S.A. de C.V.", tipo_persona="MORAL")
        db.add(titular)
        db.commit()
        titular_id = titular.id
    finally:
        db.close()
    exp_id = client.post("/api/v1/expedientes", json={"titular_id": titular_id}).json()["id"]
    r = client.post("/api/v1/documentos/upload", data={"expediente_id": exp_id, "tipo_documento": "CSF_SAT"},
                    files={"file": ("csf.pdf", b"%PDF-1.4 barrido " + os.urandom(8), "application/pdf")})
    assert r.status_code == 201, r.text
    return r.json()


def _ruta(doc: dict) -> str:
    return os.path.join(BOVEDA, f"{doc['hash_sha256'][:12]}_csf.pdf")


def _barrer() -> integridad.ResultadoBarrido:
    db = SessionLocal()
    try:
        return integridad.barrer(db)
    finally:
        db.close()


def _eventos(evento: str, documento_id: str | None = None) -> list[BitacoraAuditoria]:
    db = SessionLocal()
    try:
        consulta = db.query(BitacoraAuditoria).filter(BitacoraAuditoria.evento == evento)
        if documento_id:
            consulta = consulta.filter(BitacoraAuditoria.documento_id == documento_id)
        return consulta.order_by(BitacoraAuditoria.secuencia).all()
    finally:
        db.close()


def _dictamenes_rie03(documento_id: str) -> list[DictamenAuditoria]:
    db = SessionLocal()
    try:
        return db.query(DictamenAuditoria).filter(DictamenAuditoria.documento_id == documento_id,
                                                  DictamenAuditoria.origen == ORIGEN_RIE_03).all()
    finally:
        db.close()


def test_el_barrido_detecta_un_binario_alterado_que_nadie_leyo():
    alterado, intacto = _documento("BAR010101AAA"), _documento("BAR020202AAA")
    with open(_ruta(alterado), "wb") as archivo:
        archivo.write(b"%PDF-1.4 contenido manipulado sin abrirlo")

    resultado = _barrer()

    assert resultado.alterados >= 1 and resultado.revisados >= 2
    assert client.get(f"/api/v1/documentos/{alterado['id']}").json()["estado_integridad"] == "ALTERADO"
    assert client.get(f"/api/v1/documentos/{intacto['id']}").json()["estado_integridad"] == "INTEGRO"
    [dictamen] = _dictamenes_rie03(alterado["id"])
    assert dictamen.estatus == "NO_CONFORME"
    [alerta] = _eventos("ALERTA_INTEGRIDAD", alterado["id"])
    assert "barrido programado" in alerta.detalle["motivo"] and alerta.detalle["hash_actual"] != alterado["hash_sha256"]
    assert alerta.usuario == "SISTEMA" and alerta.hash_evento, "Registrado por el sistema y sellado"


def test_el_barrido_detecta_un_binario_eliminado():
    doc = _documento("BAR030303AAA")
    os.remove(_ruta(doc))

    _barrer()

    assert client.get(f"/api/v1/documentos/{doc['id']}").json()["estado_integridad"] == "ALTERADO"
    [alerta] = _eventos("ALERTA_INTEGRIDAD", doc["id"])
    assert alerta.detalle["hash_actual"] is None and "no se encuentra" in alerta.detalle["motivo"]


def test_cada_barrido_deja_una_constancia_y_no_repite_lo_ya_alterado():
    doc = _documento("BAR040404AAA")
    with open(_ruta(doc), "wb") as archivo:
        archivo.write(b"%PDF-1.4 otra manipulacion")
    antes = len(_eventos("VERIFICACION_INTEGRIDAD"))

    primero, segundo = _barrer(), _barrer()

    constancias = _eventos("VERIFICACION_INTEGRIDAD")
    assert len(constancias) == antes + 2, "Exactamente una constancia por barrido"
    ultima = constancias[-1]
    assert ultima.detalle["origen"] == "BARRIDO_PROGRAMADO" and ultima.hash_evento
    assert ultima.detalle["documentos_revisados"] == segundo.revisados
    assert primero.alterados >= 1 and segundo.alterados == 0, "Un documento ya ALTERADO no se vuelve a registrar"
    assert len(_eventos("ALERTA_INTEGRIDAD", doc["id"])) == 1
    assert len(_dictamenes_rie03(doc["id"])) == 1


def test_la_reevaluacion_programada_barre_y_se_puede_desactivar(monkeypatch):
    doc = _documento("BAR050505AAA")
    with open(_ruta(doc), "wb") as archivo:
        archivo.write(b"%PDF-1.4 manipulado antes del ciclo")

    monkeypatch.setenv("BARRIDO_INTEGRIDAD", "false")
    sin_barrido = reevaluacion_programada.reevaluar_todos()
    assert sin_barrido.barrido is None
    assert client.get(f"/api/v1/documentos/{doc['id']}").json()["estado_integridad"] == "INTEGRO"

    monkeypatch.delenv("BARRIDO_INTEGRIDAD")
    con_barrido = reevaluacion_programada.reevaluar_todos()
    assert con_barrido.barrido is not None and con_barrido.barrido.alterados >= 1
    assert client.get(f"/api/v1/documentos/{doc['id']}").json()["estado_integridad"] == "ALTERADO"
    expediente = client.get(f"/api/v1/expedientes/{doc['expediente_id']}").json()
    assert expediente["estado_cumplimiento"] == "NO_CONFORME", "El semáforo refleja la alteración en el mismo ciclo"


def test_la_cadena_sigue_integra_tras_los_barridos():
    assert client.get("/api/v1/bitacora/verificar-cadena").json()["integra"]


# ── PLUS-15: último sello de la cadena en los reportes (HU-45) ─────────


def _ultima_exportacion() -> BitacoraAuditoria:
    return _eventos("EXPORTACION_REPORTE")[-1]


def test_el_reporte_del_expediente_imprime_el_ultimo_sello_antes_de_exportar():
    doc = _documento("ANC010101AAA")

    xlsx = client.get(f"/api/v1/reportes/expedientes/{doc['expediente_id']}", params={"formato": "xlsx"})
    exportacion_xlsx = _ultima_exportacion()
    pdf = client.get(f"/api/v1/reportes/expedientes/{doc['expediente_id']}", params={"formato": "pdf"})
    exportacion_pdf = _ultima_exportacion()

    assert xlsx.status_code == 200 and pdf.status_code == 200
    ficha = {f[0]: f[1] for f in load_workbook(BytesIO(xlsx.content))["Dictamen"].iter_rows(values_only=True) if f[0]}
    # El sello impreso es el último eslabón antes de la exportación: el hash_anterior del evento que la registra.
    assert ficha["Último sello (SHA-256)"] == exportacion_xlsx.hash_anterior
    assert ficha["Último eslabón"] == f"Secuencia {exportacion_xlsx.secuencia - 1}"
    assert "ÍNTEGRA" in ficha["Cadena de custodia"]
    texto = "".join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages).replace("\n", "")
    assert "Constancia de la cadena de custodia" in texto and exportacion_pdf.hash_anterior in texto


def test_el_reporte_global_imprime_el_ultimo_sello_antes_de_exportar():
    xlsx = client.get("/api/v1/reportes/cumplimiento", params={"formato": "xlsx"})
    exportacion = _ultima_exportacion()
    pdf = client.get("/api/v1/reportes/cumplimiento", params={"formato": "pdf"})

    assert xlsx.status_code == 200 and pdf.status_code == 200
    ficha = {f[0]: f[1] for f in load_workbook(BytesIO(xlsx.content))["Resumen"].iter_rows(values_only=True) if f[0]}
    assert ficha["Último sello (SHA-256)"] == exportacion.hash_anterior
    texto = "".join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages).replace("\n", "")
    assert _ultima_exportacion().hash_anterior in texto
