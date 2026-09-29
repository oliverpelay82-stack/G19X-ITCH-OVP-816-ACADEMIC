"""
Reportes de auditoría exportables (Módulo 6: RF-22, RF-23, RF-24; HU-28, HU-29).

  * **Dictamen de cumplimiento por expediente (RF-23).** Consolida el semáforo,
    los documentos exigidos por las reglas, la evidencia de cada dictamen
    (reglas aplicadas, confianza, huella SHA-256, ratificación humana), las
    alertas con su recomendación y la trazabilidad de la bitácora.
  * **Reporte global de cumplimiento (RF-22).** Porcentaje global, expedientes
    conformes y no conformes, y distribución de alertas por severidad y tipo.

Ambos se emiten en PDF (presentación formal) y en Excel (.xlsx, análisis y
seguimiento), como pide RF-24. El módulo solo da formato: recibe los datos ya
evaluados por ``motor/cumplimiento.py`` y no toca la base de datos.
"""
from __future__ import annotations

import io
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from config.catalogo import etiqueta_documento, etiqueta_expediente
from models.bitacora import BitacoraAuditoria
from models.dictamen import DictamenAuditoria
from models.expediente import Expediente
from motor.cumplimiento import UNIDADES_SIN_CODIGO, EvaluacionExpediente
from motor.redaccion import texto_llano

FORMATOS = {"pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

ESTADO_LEGIBLE = {
    "CONFORME": "Conforme",
    "EN_REVISION": "En revisión",
    "NO_CONFORME": "No conforme",
    "SIN_DICTAMEN": "Sin auditar",
    "FALTANTE": "Faltante",
}
VIGENCIA_LEGIBLE = {"VIGENTE": "Vigente", "PROXIMO_A_VENCER": "Próximo a vencer", "VENCIDO": "Vencido"}
SEVERIDAD_LEGIBLE = {"CRITICA": "Crítica", "ALTA": "Alta", "MEDIA_ALTA": "Media-Alta", "MEDIA": "Media", "BAJA": "Baja"}

# Lenguaje llano en los reportes (RNF-08): los códigos internos se nombran. El
# código original se conserva donde es evidencia (la hoja de eventos de la
# ficha de custodia lleva el código sellado junto a su nombre).
EVENTO_LEGIBLE = {
    "APERTURA_EXPEDIENTE": "Apertura del expediente",
    "INGESTA_DOCUMENTO": "Documento custodiado",
    "EVALUACION_REGLAS": "Cambio del semáforo",
    "DICTAMEN_IA": "Dictamen de la IA",
    "RATIFICACION_HUMANA": "Ratificación de un analista",
    "ALERTA_INTEGRIDAD": "Integridad comprometida",
    "ARCHIVO_PROTEGIDO": "PDF protegido con contraseña",
    "EXPORTACION_REPORTE": "Reporte exportado",
    "SELLO_GENESIS": "Inicio del sellado de la bitácora",
}
ORIGEN_LEGIBLE = {
    "NIVEL_4_IA": "Análisis con IA",
    "RIE-03": "Verificación de integridad",
    "NIVEL_1": "Validación de la carga",
}
RESULTADO_LEGIBLE = {
    "CUMPLE": "cumple", "NO_CUMPLE": "no cumple", "NO_LOCALIZADO": "no localizado",
    "NO_VERIFICABLE": "no verificado", "NO_APLICA": "no aplica", "VIGENTE": "vigente",
    "PROXIMO_A_VENCER": "próximo a vencer", "VENCIDO": "vencido", "INDETERMINADO": "no verificable",
}
INTEGRIDAD_LEGIBLE = {"INTEGRO": "Íntegro", "ALTERADO": "Alterado"}
FUENTE_LEGIBLE = {
    "rules-service": "Catálogo de reglas vigente",
    "respaldo-doc02": "Inventario base de documentos (el catálogo de reglas no respondió)",
}

# Paleta del sistema de diseño (neutros stone y colores de estado).
TINTA = colors.HexColor("#1C1917")
SECUNDARIO = colors.HexColor("#57534E")
LINEA = colors.HexColor("#E7E5E4")
FONDO_SUTIL = colors.HexColor("#F5F5F4")
COLOR_ESTADO = {
    "CONFORME": (colors.HexColor("#15803D"), colors.HexColor("#F0FDF4")),
    "EN_REVISION": (colors.HexColor("#B45309"), colors.HexColor("#FFFBEB")),
    "NO_CONFORME": (colors.HexColor("#B91C1C"), colors.HexColor("#FEF2F2")),
}


@dataclass
class DatosExpediente:
    expediente: Expediente
    evaluacion: EvaluacionExpediente
    dictamenes: dict[str, DictamenAuditoria] = field(default_factory=dict)   # por documento_id
    eventos: list[BitacoraAuditoria] = field(default_factory=list)


def _fecha(valor: Any) -> str:
    if valor is None:
        return "—"
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    return valor.isoformat() if hasattr(valor, "isoformat") else str(valor)


def _ahora() -> str:
    """Hora local del servidor con su desfase, la misma referencia que las fechas de la base."""
    ahora = datetime.now().astimezone()
    desfase = ahora.strftime("%z")
    return f"{ahora.strftime('%Y-%m-%d %H:%M')} (UTC{desfase[:3]}:{desfase[3:]})"


def _pct(valor: float) -> str:
    return f"{valor:.1f} %"


RESULTADO_INDICE = {"CUMPLIDA": "Cumplida", "INCUMPLIDA": "Incumplida", "PENDIENTE": "Pendiente"}


def _texto_indice(ev: EvaluacionExpediente) -> str:
    """Índice de cumplimiento del DOC-05 (PRD §10.5) en una línea."""
    indice = ev.indice_cumplimiento
    if indice is None or indice.porcentaje is None:
        return "Sin reglas aplicables"
    texto = (f"{_pct(indice.porcentaje)}: {indice.cumplidas} de {indice.aplicables} reglas cumplidas "
             f"({indice.incumplidas} incumplidas, {indice.pendientes} pendientes)")
    if indice.bloqueado:
        texto += ". Bloqueado: " + " ".join(indice.bloqueos)
    return texto


def _filas_reglas(ev: EvaluacionExpediente) -> list[dict[str, Any]]:
    """Detalle del índice: la evidencia de cómo se calculó (hoja «Reglas»)."""
    if ev.indice_cumplimiento is None:
        return []
    return [{
        "documento": etiqueta_documento(r.tipo_documento) if r.tipo_documento else "Expediente",
        "regla": "—" if r.regla in UNIDADES_SIN_CODIGO else r.regla, "evaluacion": r.nombre, "resultado": RESULTADO_INDICE.get(r.resultado, r.resultado),
        "bloqueante": "Sí" if r.bloqueante else "",
    } for r in ev.indice_cumplimiento.reglas]


def _filas_documentos(ev: EvaluacionExpediente) -> list[dict[str, Any]]:
    """Documentos exigidos (con su archivo vigente) más los opcionales cargados."""
    por_tipo = {d.tipo_documento: d for d in ev.documentos}
    filas = []
    for o in ev.obligatorios:
        doc = por_tipo.get(o["tipo_documento"])
        filas.append({
            "tipo": o["tipo_documento"],
            "documento": etiqueta_documento(o["tipo_documento"]),
            "requisito": (f"Anexo de {etiqueta_documento(o['documento_principal'])}" if o["documento_principal"]
                          else "Obligatorio"),
            "regla": o["codigo_regla"],
            "severidad": SEVERIDAD_LEGIBLE.get(o["severidad"], o["severidad"]),
            "archivo": doc.nombre_archivo_original if doc else "No cargado",
            "estado": ESTADO_LEGIBLE.get(doc.estado if doc else "FALTANTE", "—"),
            "vigencia": _texto_vigencia(doc),
            "documento_id": doc.documento_id if doc else None,
        })
    for d in ev.documentos:
        if not d.obligatorio:
            filas.append({
                "tipo": d.tipo_documento, "documento": etiqueta_documento(d.tipo_documento), "requisito": "Opcional",
                "regla": "—", "severidad": "—", "archivo": d.nombre_archivo_original,
                "estado": ESTADO_LEGIBLE.get(d.estado, d.estado), "vigencia": _texto_vigencia(d),
                "documento_id": d.documento_id,
            })
    return filas


def _texto_vigencia(doc) -> str:
    if doc is None or not doc.vigencia:
        return "—"
    return f"{VIGENCIA_LEGIBLE.get(doc.vigencia, doc.vigencia)} ({_fecha(doc.fecha_caducidad)})"


def _evidencia(datos: DatosExpediente) -> list[dict[str, Any]]:
    """Evidencia de cada dictamen vigente (RF-23): qué se evaluó y con qué resultado."""
    filas = []
    for d in datos.evaluacion.documentos:
        dictamen = datos.dictamenes.get(d.documento_id)
        if dictamen is None:
            continue
        reglas = ", ".join(
            f"{r.get('codigo_regla')}: {RESULTADO_LEGIBLE.get(r.get('resultado'), r.get('resultado'))}"
            for r in (dictamen.reglas_aplicadas or [])
        )
        ratificacion = (
            f"{dictamen.revisado_por} ({_fecha(dictamen.revisado_en)}): {dictamen.justificacion_humana}"
            if dictamen.revisado_por_humano else "—"
        )
        filas.append({
            "documento": etiqueta_documento(d.tipo_documento),
            "archivo": d.nombre_archivo_original,
            "estatus": ESTADO_LEGIBLE.get(dictamen.estatus, dictamen.estatus),
            "preliminar": ESTADO_LEGIBLE.get(dictamen.estatus_preliminar, dictamen.estatus_preliminar),
            "origen": ORIGEN_LEGIBLE.get(dictamen.origen, dictamen.origen),
            "confianza": f"{dictamen.nivel_confianza * 100:.0f} %" if dictamen.nivel_confianza is not None else "—",
            "reglas": reglas or "—",
            "hash": dictamen.hash_sha256,
            "modelo": dictamen.modelo_utilizado or "—",
            "fecha": _fecha(dictamen.creado_en),
            "ratificacion": ratificacion,
            "observaciones": " ".join(texto_llano(o) for o in dictamen.observaciones or [] if o),
        })
    return filas


# ═══════════════════════════ Excel ═══════════════════════════════
_BORDE = Border(bottom=Side(style="thin", color="E7E5E4"))
_CABECERA = Font(bold=True, color="1C1917")
_FONDO_CABECERA = PatternFill("solid", fgColor="F5F5F4")


def _hoja_tabla(hoja, columnas: list[tuple[str, str, int]], filas: list[dict[str, Any]], fila_inicio: int = 1) -> None:
    """Escribe una tabla con cabecera fija y filtro. columnas: (clave, título, ancho)."""
    for i, (_, titulo, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_inicio, column=i, value=titulo)
        celda.font = _CABECERA
        celda.fill = _FONDO_CABECERA
        celda.border = _BORDE
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for r, fila in enumerate(filas, start=fila_inicio + 1):
        for c, (clave, _, _) in enumerate(columnas, start=1):
            celda = hoja.cell(row=r, column=c, value=fila.get(clave))
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    hoja.freeze_panes = hoja.cell(row=fila_inicio + 1, column=1)
    if filas:
        hoja.auto_filter.ref = f"A{fila_inicio}:{get_column_letter(len(columnas))}{fila_inicio + len(filas)}"


def _hoja_ficha(hoja, titulo: str, pares: list[tuple[str, Any]]) -> None:
    hoja["A1"] = titulo
    hoja["A1"].font = Font(bold=True, size=14)
    for i, (clave, valor) in enumerate(pares, start=3):
        hoja.cell(row=i, column=1, value=clave).font = Font(bold=True, color="57534E")
        hoja.cell(row=i, column=2, value=valor).alignment = Alignment(wrap_text=True, vertical="top")
    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 90


def _ficha_expediente(datos: DatosExpediente) -> list[tuple[str, Any]]:
    exp, ev = datos.expediente, datos.evaluacion
    titular = exp.titular
    return [
        ("Folio", exp.codigo_expediente),
        ("Nombre", exp.nombre or "—"),
        ("Tipo de expediente", etiqueta_expediente(exp.tipo_expediente)),
        ("Titular", titular.razon_social if titular else "—"),
        ("RFC", titular.rfc if titular else "—"),
        ("Fecha de apertura", _fecha(exp.creado_en)),
        ("Estado de cumplimiento", ESTADO_LEGIBLE.get(ev.estado, ev.estado)),
        ("Índice de cumplimiento", _texto_indice(ev)),
        ("Documentos obligatorios presentes", _pct(ev.porcentaje_completitud)),
        ("Documentos faltantes", ", ".join(etiqueta_documento(t) for t in ev.faltantes) or "Ninguno"),
        ("Alertas vigentes", len(ev.alertas)),
        ("Criterio de obligatoriedad", FUENTE_LEGIBLE.get(ev.fuente_obligatorios, ev.fuente_obligatorios)
         + (f" — {ev.advertencia}" if ev.advertencia else "")),
        ("Reporte emitido", _ahora()),
    ]


def expediente_xlsx(datos: DatosExpediente) -> bytes:
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Dictamen"
    _hoja_ficha(resumen, "Dictamen de Cumplimiento Documental", _ficha_expediente(datos))
    fila = len(_ficha_expediente(datos)) + 5
    resumen.cell(row=fila, column=1, value="Motivos del semáforo").font = Font(bold=True)
    for i, motivo in enumerate(datos.evaluacion.motivos, start=fila + 1):
        resumen.cell(row=i, column=2, value=motivo).alignment = Alignment(wrap_text=True)

    _hoja_tabla(libro.create_sheet("Documentos"), [
        ("documento", "Documento", 38), ("requisito", "Requisito", 30), ("regla", "Regla", 16),
        ("severidad", "Severidad", 12), ("archivo", "Archivo vigente", 40), ("estado", "Estado", 14),
        ("vigencia", "Vigencia", 28),
    ], _filas_documentos(datos.evaluacion))

    _hoja_tabla(libro.create_sheet("Reglas"), [
        ("documento", "Documento", 34), ("regla", "Regla", 18), ("evaluacion", "Evaluación", 60),
        ("resultado", "Resultado", 13), ("bloqueante", "Bloquea", 9),
    ], _filas_reglas(datos.evaluacion))

    _hoja_tabla(libro.create_sheet("Evidencia"), [
        ("documento", "Documento", 32), ("archivo", "Archivo", 34), ("estatus", "Estatus", 14),
        ("preliminar", "Estatus preliminar", 16), ("origen", "Origen", 12), ("confianza", "Confianza", 11),
        ("reglas", "Reglas aplicadas", 40), ("hash", "SHA-256", 68), ("modelo", "Modelo", 20),
        ("fecha", "Fecha", 17), ("ratificacion", "Ratificación humana", 50), ("observaciones", "Observaciones", 70),
    ], _evidencia(datos))

    _hoja_tabla(libro.create_sheet("Alertas"), [
        ("severidad", "Severidad", 12), ("titulo", "Alerta", 42), ("detalle", "Detalle", 60),
        ("recomendacion", "Recomendación", 70), ("regla", "Regla", 14),
    ], [{**a.__dict__, "severidad": SEVERIDAD_LEGIBLE.get(a.severidad, a.severidad)} for a in datos.evaluacion.alertas])

    _hoja_tabla(libro.create_sheet("Bitácora"), [
        ("fecha", "Fecha", 18), ("evento", "Evento", 30), ("usuario", "Usuario", 24),
        ("ip", "IP de origen", 16), ("detalle", "Detalle registrado", 90),
    ], [{"fecha": _fecha(e.creado_en), "evento": EVENTO_LEGIBLE.get(e.evento, e.evento), "usuario": e.usuario,
         "ip": e.ip_origen or "—", "detalle": _resumen_detalle(e.detalle)} for e in datos.eventos])
    return _guardar(libro)


def global_xlsx(evaluados: list[tuple[Expediente, EvaluacionExpediente]]) -> bytes:
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    indicadores = _indicadores(evaluados)
    _hoja_ficha(resumen, "Reporte de Cumplimiento Documental", [
        ("Expedientes evaluados", indicadores["total"]),
        ("Expedientes conformes", _pct(indicadores["porcentaje_global"])),
        ("Índice de cumplimiento promedio", _pct(indicadores["indice_promedio"])
         if indicadores["indice_promedio"] is not None else "Sin reglas aplicables"),
        ("Conformes", indicadores["conformes"]),
        ("En revisión", indicadores["en_revision"]),
        ("No conformes", indicadores["no_conformes"]),
        ("Documentos obligatorios presentes (promedio)", _pct(indicadores["completitud_promedio"])),
        ("Alertas vigentes", indicadores["alertas"]),
        *[(f"Alertas {SEVERIDAD_LEGIBLE.get(s, s).lower()}", n) for s, n in indicadores["alertas_por_severidad"]],
        *[(f"Expedientes {etiqueta_expediente(t).lower()}s", n) for t, n in indicadores["por_tipo"]],
        ("Reporte emitido", _ahora()),
    ])
    _hoja_tabla(libro.create_sheet("Expedientes"), [
        ("folio", "Folio", 26), ("nombre", "Nombre", 30), ("tipo", "Tipo", 14), ("titular", "Titular", 38),
        ("rfc", "RFC", 16), ("estado", "Estado", 14), ("indice", "Índice de cumplimiento %", 14),
        ("completitud", "Documentos obligatorios %", 14),
        ("faltantes", "Faltantes", 50), ("alertas", "Alertas", 9), ("apertura", "Apertura", 17),
    ], _filas_expedientes(evaluados))
    _hoja_tabla(libro.create_sheet("Alertas"), [
        ("folio", "Folio", 26), ("titular", "Titular", 34), ("severidad", "Severidad", 12),
        ("titulo", "Alerta", 42), ("recomendacion", "Recomendación", 80),
    ], _filas_alertas(evaluados))
    return _guardar(libro)


def _guardar(libro: Workbook) -> bytes:
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _resumen_detalle(detalle: dict[str, Any] | None) -> str:
    if not detalle:
        return ""
    partes = []
    for clave, valor in detalle.items():
        if isinstance(valor, (list, dict)):
            if not valor:
                continue
            valor = "; ".join(str(v.get("titulo", v)) if isinstance(v, dict) else str(v) for v in valor) \
                if isinstance(valor, list) else str(valor)
        partes.append(f"{clave}: {valor}")
    return " · ".join(partes)[:1000]


# ═══════════════════════ Indicadores (RF-22) ═══════════════════════
def _indicadores(evaluados: list[tuple[Expediente, EvaluacionExpediente]]) -> dict[str, Any]:
    estados = Counter(ev.estado for _, ev in evaluados)
    total = len(evaluados)
    indices = [ev.indice_cumplimiento.porcentaje for _, ev in evaluados
               if ev.indice_cumplimiento and ev.indice_cumplimiento.porcentaje is not None]
    alertas = [a for _, ev in evaluados for a in ev.alertas]
    orden = list(SEVERIDAD_LEGIBLE)
    por_severidad = sorted(Counter(a.severidad for a in alertas).items(),
                           key=lambda p: orden.index(p[0]) if p[0] in orden else 99)
    return {
        "total": total,
        "conformes": estados.get("CONFORME", 0),
        "en_revision": estados.get("EN_REVISION", 0),
        "no_conformes": estados.get("NO_CONFORME", 0),
        # Proporción de expedientes que cumplen al 100 % (RF-22).
        "porcentaje_global": round(estados.get("CONFORME", 0) / total * 100, 1) if total else 0.0,
        # Promedio del índice del DOC-05: cada expediente pesa igual (PRD §10.5).
        "indice_promedio": round(sum(indices) / len(indices), 1) if indices else None,
        "completitud_promedio": round(sum(ev.porcentaje_completitud for _, ev in evaluados) / total, 1) if total else 0.0,
        "alertas": len(alertas),
        "alertas_por_severidad": por_severidad,
        "por_tipo": sorted(Counter(exp.tipo_expediente for exp, _ in evaluados).items()),
    }


def _filas_expedientes(evaluados) -> list[dict[str, Any]]:
    return [{
        "folio": exp.codigo_expediente,
        "nombre": exp.nombre or "—",
        "tipo": etiqueta_expediente(exp.tipo_expediente),
        "titular": exp.titular.razon_social if exp.titular else "—",
        "rfc": exp.titular.rfc if exp.titular else "—",
        "estado": ESTADO_LEGIBLE.get(ev.estado, ev.estado),
        "completitud": ev.porcentaje_completitud,
        "indice": ev.indice_cumplimiento.porcentaje if ev.indice_cumplimiento else None,
        "faltantes": ", ".join(etiqueta_documento(t) for t in ev.faltantes) or "—",
        "alertas": len(ev.alertas),
        "apertura": _fecha(exp.creado_en),
    } for exp, ev in evaluados]


def _filas_alertas(evaluados) -> list[dict[str, Any]]:
    return [{
        "folio": exp.codigo_expediente,
        "titular": exp.titular.razon_social if exp.titular else "—",
        "severidad": SEVERIDAD_LEGIBLE.get(a.severidad, a.severidad),
        "titulo": a.titulo,
        "recomendacion": a.recomendacion,
    } for exp, ev in evaluados for a in ev.alertas]


# ═══════════════════════════ PDF ═══════════════════════════════
def _estilos() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("titulo", parent=base["Title"], fontName="Helvetica-Bold", fontSize=16,
                                 textColor=TINTA, alignment=TA_LEFT, spaceAfter=2),
        "subtitulo": ParagraphStyle("subtitulo", parent=base["Normal"], fontSize=9, textColor=SECUNDARIO,
                                    spaceAfter=10),
        "seccion": ParagraphStyle("seccion", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=10.5,
                                  textColor=TINTA, spaceBefore=12, spaceAfter=5),
        "texto": ParagraphStyle("texto", parent=base["Normal"], fontSize=8.5, leading=11, textColor=TINTA),
        "celda": ParagraphStyle("celda", parent=base["Normal"], fontSize=7.5, leading=9.5, textColor=TINTA),
        "celda_mono": ParagraphStyle("celda_mono", parent=base["Normal"], fontName="Courier", fontSize=6.5,
                                     leading=8.5, textColor=SECUNDARIO),
        "cabecera": ParagraphStyle("cabecera", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7.5,
                                   leading=9.5, textColor=SECUNDARIO),
    }


def _esc(texto: Any) -> str:
    return str(texto if texto is not None else "—").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _tabla(filas: list[list[Any]], anchos: list[float], estilos, cabecera: list[str]) -> Table:
    datos = [[Paragraph(_esc(c), estilos["cabecera"]) for c in cabecera]]
    for fila in filas:
        datos.append([c if isinstance(c, Paragraph) else Paragraph(_esc(c), estilos["celda"]) for c in fila])
    tabla = Table(datos, colWidths=anchos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), FONDO_SUTIL),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINEA),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    return tabla


def _ficha(pares: list[tuple[str, Any]], estilos, ancho: float) -> Table:
    filas = [[Paragraph(_esc(k), estilos["cabecera"]), Paragraph(_esc(v), estilos["texto"])] for k, v in pares]
    tabla = Table(filas, colWidths=[ancho * 0.3, ancho * 0.7])
    tabla.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINEA),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return tabla


def _sello_estado(estado: str, texto: str, estilos, ancho: float) -> Table:
    """Semáforo del expediente: color de estado siempre acompañado de su etiqueta."""
    tinta, fondo = COLOR_ESTADO.get(estado, (SECUNDARIO, FONDO_SUTIL))
    estilo = ParagraphStyle("sello", parent=estilos["texto"], fontName="Helvetica-Bold", fontSize=11, textColor=tinta)
    tabla = Table([[Paragraph(_esc(texto), estilo)]], colWidths=[ancho])
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), fondo),
        ("BOX", (0, 0), (-1, -1), 0.6, tinta),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    return tabla


def _pie(lienzo, doc) -> None:
    lienzo.saveState()
    lienzo.setFont("Helvetica", 7)
    lienzo.setFillColor(SECUNDARIO)
    lienzo.drawString(doc.leftMargin, 10 * mm,
                      "PluriOne S.A. de C.V. · Motor de Validación de Cumplimiento Documental · Documento generado "
                      "automáticamente a partir de la bitácora inmutable")
    lienzo.drawRightString(doc.pagesize[0] - doc.rightMargin, 10 * mm, f"Página {doc.page}")
    lienzo.restoreState()


def _documento_pdf(historia: list, titulo: str) -> bytes:
    salida = io.BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=letter, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=titulo, author="PluriOne — Motor de Cumplimiento")
    doc.build(historia, onFirstPage=_pie, onLaterPages=_pie)
    return salida.getvalue()


def expediente_pdf(datos: DatosExpediente) -> bytes:
    e = _estilos()
    ancho = letter[0] - 32 * mm
    exp, ev = datos.expediente, datos.evaluacion
    estado = ESTADO_LEGIBLE.get(ev.estado, ev.estado)
    porcentaje = ev.indice_cumplimiento.porcentaje if ev.indice_cumplimiento else None
    indice = _pct(porcentaje) if porcentaje is not None else "—"
    historia: list = [
        Paragraph("Dictamen de Cumplimiento Documental", e["titulo"]),
        Paragraph(f"Expediente {_esc(exp.codigo_expediente)} · emitido el {_ahora()}", e["subtitulo"]),
        _sello_estado(ev.estado, f"{estado} · índice de cumplimiento {indice} · documentos obligatorios "
                                 f"{_pct(ev.porcentaje_completitud)} · {len(ev.alertas)} alertas vigentes", e, ancho),
        Spacer(1, 8),
        _ficha(_ficha_expediente(datos)[:9], e, ancho),
        Paragraph("Motivos del semáforo", e["seccion"]),
        *[Paragraph(f"• {_esc(texto_llano(m))}", e["texto"]) for m in ev.motivos],
        Paragraph("Documentos exigidos por el catálogo de reglas", e["seccion"]),
        _tabla(
            [[f["documento"], f["requisito"], f["regla"], f["archivo"], f["estado"], f["vigencia"]]
             for f in _filas_documentos(ev)],
            [ancho * p for p in (0.22, 0.15, 0.13, 0.22, 0.11, 0.17)], e,
            ["Documento", "Requisito", "Regla", "Archivo vigente", "Estado", "Vigencia"],
        ),
    ]

    if ev.alertas:
        historia += [
            Paragraph("Alertas y recomendaciones", e["seccion"]),
            _tabla(
                [[SEVERIDAD_LEGIBLE.get(a.severidad, a.severidad), f"{a.titulo}. {a.detalle}", a.recomendacion]
                 for a in ev.alertas],
                [ancho * p for p in (0.11, 0.42, 0.47)], e, ["Severidad", "Hallazgo", "Recomendación"],
            ),
        ]

    evidencia = _evidencia(datos)
    if evidencia:
        historia.append(Paragraph("Evidencia de auditoría por documento", e["seccion"]))
        for f in evidencia:
            historia.append(KeepTogether([
                _ficha([
                    ("Documento", f"{f['documento']} — {f['archivo']}"),
                    ("Dictamen", f"{f['estatus']} (preliminar: {f['preliminar']}, origen: {f['origen']}, "
                                 f"confianza {f['confianza']}, {f['fecha']})"),
                    ("Reglas aplicadas", f["reglas"]),
                    ("Ratificación humana", f["ratificacion"]),
                    ("Observaciones", f["observaciones"] or "—"),
                ], e, ancho),
                Paragraph(f"SHA-256: {_esc(f['hash'])}", e["celda_mono"]),
                Spacer(1, 6),
            ]))

    if datos.eventos:
        historia += [
            Paragraph("Trazabilidad (bitácora inmutable)", e["seccion"]),
            _tabla(
                [[_fecha(ev_.creado_en), EVENTO_LEGIBLE.get(ev_.evento, ev_.evento), ev_.usuario, ev_.ip_origen or "—"]
                 for ev_ in datos.eventos[:40]],
                [ancho * p for p in (0.2, 0.3, 0.3, 0.2)], e, ["Fecha", "Evento", "Usuario", "IP de origen"],
            ),
        ]
    return _documento_pdf(historia, f"Dictamen {exp.codigo_expediente}")


def global_pdf(evaluados: list[tuple[Expediente, EvaluacionExpediente]]) -> bytes:
    e = _estilos()
    ancho = letter[0] - 32 * mm
    ind = _indicadores(evaluados)
    historia: list = [
        Paragraph("Reporte de Cumplimiento Documental", e["titulo"]),
        Paragraph(f"{ind['total']} expedientes evaluados · emitido el {_ahora()}", e["subtitulo"]),
        _ficha([
            ("Expedientes conformes", f"{_pct(ind['porcentaje_global'])} de los expedientes"),
            ("Índice de cumplimiento promedio", _pct(ind["indice_promedio"])
             if ind["indice_promedio"] is not None else "Sin reglas aplicables"),
            ("Conformes / En revisión / No conformes",
             f"{ind['conformes']} / {ind['en_revision']} / {ind['no_conformes']}"),
            ("Documentos obligatorios presentes (promedio)", _pct(ind["completitud_promedio"])),
            ("Alertas vigentes", ", ".join(f"{SEVERIDAD_LEGIBLE.get(s, s)}: {n}" for s, n in
                                           ind["alertas_por_severidad"]) or "Ninguna"),
            ("Expedientes por tipo", ", ".join(f"{etiqueta_expediente(t)}: {n}" for t, n in ind["por_tipo"]) or "—"),
        ], e, ancho),
        Paragraph("Expedientes", e["seccion"]),
        _tabla(
            [[f["folio"], f["tipo"], f["titular"], f["estado"],
              f"{f['indice']:.1f} %" if f["indice"] is not None else "—", f"{f['completitud']:.1f} %", f["alertas"]]
             for f in _filas_expedientes(evaluados)],
            [ancho * p for p in (0.22, 0.11, 0.29, 0.12, 0.1, 0.09, 0.07)], e,
            ["Folio", "Tipo", "Titular", "Estado", "Índice", "Obligatorios", "Alertas"],
        ),
    ]
    alertas = _filas_alertas(evaluados)
    if alertas:
        historia += [
            Paragraph("Alertas y recomendaciones", e["seccion"]),
            _tabla([[a["folio"], a["severidad"], a["titulo"], a["recomendacion"]] for a in alertas],
                   [ancho * p for p in (0.2, 0.1, 0.28, 0.42)], e, ["Folio", "Severidad", "Alerta", "Recomendación"]),
        ]
    return _documento_pdf(historia, "Reporte de Cumplimiento Documental")



# ═══════════════════ Ficha de custodia (P3, PLUS-06) ═══════════════════
LEYENDA_CUSTODIA = (
    "Ficha alineada con ISO 30301 (integridad, metadatos y trazabilidad del documento) y con NOM-151-SCFI-2016 "
    "(conservación íntegra de mensajes de datos). No sustituye la constancia de conservación que emite un "
    "Prestador de Servicios de Certificación acreditado."
)


@dataclass
class DatosCustodia:
    """Ciclo de vida de un documento para una auditoría externa (P3, PLUS-06)."""
    documento: Any                                   # models.documento.Documento
    expediente: Expediente
    integro: bool                                    # la huella recalculada coincide con la de la carga (RIE-03)
    hash_actual: str | None                          # huella recalculada; None si el binario no está en la bóveda
    versiones: list[Any] = field(default_factory=list)          # documentos del mismo tipo en el expediente
    dictamenes: list[DictamenAuditoria] = field(default_factory=list)
    eventos: list[BitacoraAuditoria] = field(default_factory=list)
    cadena: dict[str, Any] = field(default_factory=dict)       # motor.bitacora.verificar_cadena


def _ficha_custodia(datos: DatosCustodia) -> list[tuple[str, Any]]:
    doc, exp = datos.documento, datos.expediente
    titular = exp.titular
    posicion = next((i for i, v in enumerate(datos.versiones, start=1) if v.id == doc.id), None)
    vigente = bool(datos.versiones) and datos.versiones[-1].id == doc.id
    estado_version = "vigente" if vigente else "sustituida por una versión posterior"
    return [
        ("Documento", doc.nombre_archivo_original),
        ("Tipo documental", etiqueta_documento(doc.tipo_documento)),
        ("Expediente", f"{exp.codigo_expediente} · {etiqueta_expediente(exp.tipo_expediente)}"),
        ("Titular", f"{titular.razon_social} ({titular.rfc})" if titular else "—"),
        ("Ingesta", _fecha(doc.creado_en)),
        ("Tamaño y formato", f"{doc.tamanio_bytes} bytes · {doc.mime_type}"),
        ("Huella SHA-256 de la carga", doc.hash_sha256),
        ("Huella recalculada al emitir", datos.hash_actual or "El binario no está en la bóveda"),
        ("Integridad", "Íntegro: las huellas coinciden" if datos.integro
         else "ALTERADO: la huella ya no coincide; el documento está bloqueado"),
        ("Versión", f"{posicion or '—'} de {len(datos.versiones)} · {estado_version}"),
        ("Ficha emitida", _ahora()),
    ]


def _texto_cadena(cadena: dict[str, Any]) -> str:
    if not cadena.get("sellada"):
        return "La bitácora aún no tiene eventos sellados."
    estado = "ÍNTEGRA" if cadena.get("integra") else f"COMPROMETIDA ({cadena.get('total_hallazgos', 0)} hallazgos)"
    return (f"Cadena de la bitácora {estado}: {cadena.get('eventos_sellados')} eventos sellados y "
            f"{cadena.get('eventos_previos')} previos al sellado, cubiertos por el génesis.")


def _filas_versiones(datos: DatosCustodia) -> list[dict[str, Any]]:
    ultimo = datos.versiones[-1].id if datos.versiones else None
    return [{
        "n": i, "archivo": v.nombre_archivo_original, "fecha": _fecha(v.creado_en), "hash": v.hash_sha256,
        "estado": "Vigente" if v.id == ultimo else "Sustituida",
        "integridad": INTEGRIDAD_LEGIBLE.get(v.estado_integridad or "INTEGRO", v.estado_integridad), "esta": "← esta ficha" if v.id == datos.documento.id else "",
    } for i, v in enumerate(datos.versiones, start=1)]


def _filas_dictamenes_custodia(datos: DatosCustodia) -> list[dict[str, Any]]:
    filas = []
    for d in datos.dictamenes:
        filas.append({
            "fecha": _fecha(d.creado_en),
            "origen": ORIGEN_LEGIBLE.get(d.origen, d.origen),
            "resultado": f"{ESTADO_LEGIBLE.get(d.estatus_preliminar, d.estatus_preliminar)} → "
                         f"{ESTADO_LEGIBLE.get(d.estatus, d.estatus)}",
            "confianza": f"{d.nivel_confianza * 100:.0f} %" if d.nivel_confianza is not None else "—",
            "huella": ("Coincide con la carga" if d.hash_sha256 == datos.documento.hash_sha256
                       else f"Distinta: {d.hash_sha256}"),
            "ratificacion": (f"{d.revisado_por} ({_fecha(d.revisado_en)}): {d.justificacion_humana}"
                             if d.revisado_por_humano else "Sin ratificar"),
            "modelo": d.modelo_utilizado or "—",
        })
    return filas


def _filas_eventos_custodia(datos: DatosCustodia) -> list[dict[str, Any]]:
    return [{
        "secuencia": e.secuencia if e.secuencia is not None else "previo",
        "fecha": _fecha(e.creado_en), "evento": EVENTO_LEGIBLE.get(e.evento, e.evento), "codigo": e.evento,
        "usuario": e.usuario, "ip": e.ip_origen or "—",
        "hash": e.hash_evento or "cubierto por el génesis", "detalle": _resumen_detalle(e.detalle),
    } for e in datos.eventos]


def ficha_custodia_xlsx(datos: DatosCustodia) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ficha"
    pares = _ficha_custodia(datos) + [("Cadena de la bitácora", _texto_cadena(datos.cadena)),
                                      ("Último sello", datos.cadena.get("ultimo_hash") or "—"),
                                      ("Algoritmo", datos.cadena.get("algoritmo") or "—"),
                                      ("Alcance normativo", LEYENDA_CUSTODIA)]
    _hoja_ficha(hoja, "Ficha de Custodia Documental", pares)
    _hoja_tabla(libro.create_sheet("Versiones"), [
        ("n", "#", 5), ("archivo", "Archivo", 40), ("fecha", "Ingesta", 17), ("hash", "SHA-256", 68),
        ("estado", "Estado", 12), ("integridad", "Integridad", 12), ("esta", "", 14),
    ], _filas_versiones(datos))
    _hoja_tabla(libro.create_sheet("Dictámenes"), [
        ("fecha", "Fecha", 17), ("origen", "Origen", 12), ("resultado", "Preliminar → vigente", 28),
        ("confianza", "Confianza", 11), ("huella", "Huella evaluada", 40), ("ratificacion", "Ratificación humana", 60),
        ("modelo", "Modelo", 20),
    ], _filas_dictamenes_custodia(datos))
    _hoja_tabla(libro.create_sheet("Eventos"), [
        ("secuencia", "Secuencia", 10), ("fecha", "Fecha", 17), ("evento", "Evento", 30),
        ("codigo", "Código sellado", 22), ("usuario", "Usuario", 22), ("ip", "IP de origen", 16),
        ("hash", "Sello del evento", 68), ("detalle", "Detalle registrado", 80),
    ], _filas_eventos_custodia(datos))
    _hoja_tabla(libro.create_sheet("Cadena"), [
        ("secuencia", "Secuencia", 10), ("evento_id", "Evento", 38), ("problema", "Hallazgo", 90),
    ], datos.cadena.get("hallazgos") or [])
    return _guardar(libro)


def ficha_custodia_pdf(datos: DatosCustodia) -> bytes:
    e = _estilos()
    ancho = letter[0] - 32 * mm
    doc = datos.documento
    integra = datos.cadena.get("integra", True)
    estado = "CONFORME" if datos.integro and integra else "NO_CONFORME"
    sello = ("Documento íntegro" if datos.integro else "Documento ALTERADO") + " · " + _texto_cadena(datos.cadena)
    historia: list = [
        Paragraph("Ficha de Custodia Documental", e["titulo"]),
        Paragraph(f"{_esc(doc.nombre_archivo_original)} · emitida el {_ahora()}", e["subtitulo"]),
        _sello_estado(estado, sello, e, ancho),
        Spacer(1, 8),
        _ficha(_ficha_custodia(datos), e, ancho),
        Paragraph("Versiones del documento en el expediente", e["seccion"]),
        _tabla([[f["n"], f["archivo"], f["fecha"], Paragraph(_esc(f["hash"]), e["celda_mono"]), f["estado"]]
                for f in _filas_versiones(datos)],
               [ancho * p for p in (0.05, 0.25, 0.15, 0.42, 0.13)], e,
               ["#", "Archivo", "Ingesta", "SHA-256", "Estado"]),
    ]
    dictamenes = _filas_dictamenes_custodia(datos)
    historia.append(Paragraph("Dictámenes y ratificaciones", e["seccion"]))
    if dictamenes:
        historia.append(_tabla(
            [[f["fecha"], f["origen"], f["resultado"], f["confianza"], f["huella"], f["ratificacion"]]
             for f in dictamenes],
            [ancho * p for p in (0.13, 0.11, 0.19, 0.09, 0.18, 0.30)], e,
            ["Fecha", "Origen", "Preliminar → vigente", "Confianza", "Huella evaluada", "Ratificación"]))
    else:
        historia.append(Paragraph("El documento no tiene dictámenes.", e["texto"]))
    historia += [
        Paragraph("Eventos de la bitácora del documento, con su sello", e["seccion"]),
        _tabla([[f["secuencia"], f["fecha"], f["evento"], f["usuario"], Paragraph(_esc(f["hash"]), e["celda_mono"])]
                for f in _filas_eventos_custodia(datos)],
               [ancho * p for p in (0.09, 0.15, 0.2, 0.16, 0.4)], e,
               ["Secuencia", "Fecha", "Evento", "Usuario", "Sello del evento"]),
        Paragraph("Verificación de la cadena de custodia", e["seccion"]),
        _ficha([
            ("Resultado", _texto_cadena(datos.cadena)),
            ("Último sello", datos.cadena.get("ultimo_hash") or "—"),
            ("Verificada el", datos.cadena.get("verificado_en") or "—"),
            ("Algoritmo", datos.cadena.get("algoritmo") or "—"),
        ], e, ancho),
    ]
    for h in (datos.cadena.get("hallazgos") or [])[:10]:
        historia.append(Paragraph(f"• Secuencia {_esc(h.get('secuencia'))}: {_esc(h.get('problema'))}", e["texto"]))
    historia += [Spacer(1, 10), Paragraph(_esc(LEYENDA_CUSTODIA), e["subtitulo"])]
    return _documento_pdf(historia, f"Ficha de custodia {doc.nombre_archivo_original}")
